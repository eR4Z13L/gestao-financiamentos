"""Migracoes da aba CLIENTES pro formato atual (endereco em campos separados,
com COMPLEMENTO e UF, mais NOME DO PAI / NOME DA MÃE / PROFISSÃO).

Aceita dois formatos de origem:
- "legado":  uma unica coluna de texto corrido "ENDEREÇO" (11 colunas)
- "intermediario": endereco ja em campos, mas sem COMPLEMENTO e com a UF
  grudada na cidade ("Fortaleza/CE") - 19 colunas

Seguranca, na ordem em que acontece (migrar_planilha):
1. so roda se a aba estiver exatamente num dos formatos antigos (e nao estiver
   bloqueada pelo Excel);
2. copia o arquivo inteiro pra um backup ao lado e confere byte a byte (hash);
3. monta o resultado num arquivo TEMPORARIO e o RELE, conferindo que todas as
   outras colunas/abas ficaram identicas ao original;
4. so entao substitui o arquivo real (troca atomica).
Se qualquer conferencia falhar, o arquivo real nunca chega a ser tocado.

Endereco que nao segue um padrao claro NAO e separado (core/endereco.py): o
texto original vai inteiro pra coluna "ENDEREÇO (REVISAR)", nada e perdido nem
adivinhado. Numa migracao do formato intermediario, esses textos sao
reavaliados pelo parser atual - mas so os de quem ainda nao tem NENHUM campo
de endereco preenchido (quem ja comecou a revisar na mao nao e tocado).
"""

from __future__ import annotations

import hashlib
import os
import shutil
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from core import data_store as bd
from core.endereco import EnderecoSeparado, separar_cidade_uf, separar_endereco

FORMATO_LEGADO = "legado"
FORMATO_INTERMEDIARIO = "intermediario"

COLUNAS_LEGADO = [
    "DATA CADASTRO",
    "CPF/CNPJ",
    "VENDEDOR",
    "TIPO",
    "CLIENTE",
    "NASCIMENTO",
    "CELULAR",
    "ENDEREÇO",
    "VINCULADO",
    "REDE SOCIAL",
    "EMAIL",
]

COLUNAS_INTERMEDIARIO = [
    "DATA CADASTRO",
    "CPF/CNPJ",
    "VENDEDOR",
    "TIPO",
    "CLIENTE",
    "NASCIMENTO",
    "CELULAR",
    "CEP",
    "LOGRADOURO",
    "NÚMERO",
    "BAIRRO",
    "CIDADE",
    "ENDEREÇO (REVISAR)",
    "VINCULADO",
    "REDE SOCIAL",
    "EMAIL",
    "NOME DO PAI",
    "NOME DA MÃE",
    "PROFISSÃO",
]

_COLUNAS_POR_FORMATO = {FORMATO_LEGADO: COLUNAS_LEGADO, FORMATO_INTERMEDIARIO: COLUNAS_INTERMEDIARIO}
_ROTULO_BACKUP = {FORMATO_LEGADO: "pre-endereco", FORMATO_INTERMEDIARIO: "pre-complemento-uf"}

_COLUNAS_ENDERECO = (
    "CEP", "LOGRADOURO", "NÚMERO", "COMPLEMENTO", "BAIRRO", "CIDADE", "UF", "ENDEREÇO (REVISAR)",
)


def _colunas_preservadas(formato: str) -> list[str]:
    """Colunas que so mudam de posicao (mesmo nome e mesmo valor antes e depois)."""
    return [c for c in _COLUNAS_POR_FORMATO[formato] if c not in _COLUNAS_ENDERECO and c != "ENDEREÇO"]


_LARGURAS = {
    "DATA CADASTRO": 13, "CPF/CNPJ": 16, "VENDEDOR": 12, "TIPO": 10, "CLIENTE": 30, "NASCIMENTO": 13,
    "CELULAR": 16, "CEP": 11, "LOGRADOURO": 34, "NÚMERO": 9, "COMPLEMENTO": 16, "BAIRRO": 22,
    "CIDADE": 24, "UF": 5, "ENDEREÇO (REVISAR)": 44, "VINCULADO": 20, "REDE SOCIAL": 22, "EMAIL": 28,
    "NOME DO PAI": 28, "NOME DA MÃE": 28, "PROFISSÃO": 22,
}


class ErroMigracao(Exception):
    """Migracao recusada ou abortada antes de tocar no arquivo real."""


@dataclass
class LinhaEndereco:
    linha_planilha: int  # numero da linha no Excel (cabecalho = 1)
    cliente: str
    original: str
    separado: EnderecoSeparado | None
    motivo: str  # "" quando separou


@dataclass
class Relatorio:
    formato: str = ""
    total_clientes: int = 0
    sem_endereco: int = 0
    ja_separados: int = 0  # (intermediario) ja tinham os campos preenchidos - nao tocados
    separados: list[LinhaEndereco] = field(default_factory=list)  # separados AGORA
    revisar: list[LinhaEndereco] = field(default_factory=list)  # continuam sem separar
    cidades_divididas: int = 0  # (intermediario) "Cidade/UF" -> CIDADE + UF
    cidades_sem_uf: list[tuple[int, str, str]] = field(default_factory=list)  # (linha, cliente, cidade)
    caminho_backup: Path | None = None


def _hash_arquivo(caminho: Path) -> str:
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


def _norm(valor) -> str:
    return bd._normalizar_texto(valor)


def _ler_registros_antigos(ws, colunas: list[str]) -> list[tuple[int, dict]]:
    """[(numero da linha no Excel, {coluna antiga: valor}), ...] - sem
    nenhuma normalizacao (o que estava na celula, como estava)."""
    registros = []
    for numero, linha in enumerate(ws.iter_rows(min_row=2, max_col=len(colunas), values_only=True), start=2):
        if all(v is None or v == "" for v in linha):
            continue
        registros.append((numero, dict(zip(colunas, linha))))
    return registros


def _campos_do_endereco(separado: EnderecoSeparado) -> dict:
    return {
        "CEP": separado.cep,
        "LOGRADOURO": separado.logradouro,
        "NÚMERO": separado.numero,
        "COMPLEMENTO": separado.complemento,
        "BAIRRO": separado.bairro,
        "CIDADE": separado.cidade,
        "UF": separado.uf,
    }


def _montar_novos_registros(antigos: list[tuple[int, dict]], formato: str) -> tuple[list[dict], Relatorio]:
    relatorio = Relatorio(formato=formato, total_clientes=len(antigos))
    preservadas = _colunas_preservadas(formato)
    novos: list[dict] = []

    for numero, antigo in antigos:
        novo = {coluna: antigo[coluna] for coluna in preservadas}
        cliente = _norm(antigo["CLIENTE"])

        if formato == FORMATO_LEGADO:
            texto = _norm(antigo["ENDEREÇO"])
            if not texto:
                relatorio.sem_endereco += 1
            else:
                _classificar(numero, cliente, texto, novo, relatorio)
        else:
            _migrar_intermediario(numero, cliente, antigo, novo, relatorio)
        novos.append(novo)
    return novos, relatorio


def _classificar(numero: int, cliente: str, texto: str, novo: dict, relatorio: Relatorio) -> None:
    """Tenta separar `texto`; preenche `novo` com os campos ou, se nao der,
    com o texto original em ENDEREÇO (REVISAR)."""
    separado, motivo = separar_endereco(texto)
    item = LinhaEndereco(numero, cliente, texto, separado, motivo)
    if separado is None:
        novo["ENDEREÇO (REVISAR)"] = texto
        relatorio.revisar.append(item)
    else:
        novo.update(_campos_do_endereco(separado))
        relatorio.separados.append(item)


def _migrar_intermediario(numero: int, cliente: str, antigo: dict, novo: dict, relatorio: Relatorio) -> None:
    for coluna in ("CEP", "LOGRADOURO", "NÚMERO", "BAIRRO", "ENDEREÇO (REVISAR)"):
        novo[coluna] = antigo[coluna]

    cidade_antiga = _norm(antigo["CIDADE"])
    cidade, uf = separar_cidade_uf(cidade_antiga)
    novo["CIDADE"], novo["UF"] = cidade, uf
    if uf:
        relatorio.cidades_divididas += 1
    elif cidade_antiga:
        relatorio.cidades_sem_uf.append((numero, cliente, cidade_antiga))

    tem_campos = any(_norm(antigo[c]) for c in ("CEP", "LOGRADOURO", "NÚMERO", "BAIRRO", "CIDADE"))
    texto_revisar = _norm(antigo["ENDEREÇO (REVISAR)"])
    if tem_campos:
        relatorio.ja_separados += 1  # inclui quem esta com revisao pela metade: nao mexe em nada
    elif texto_revisar:
        # nenhum campo preenchido ainda: o parser atual (que sabe de complemento
        # e UF) pode conseguir o que o anterior nao conseguiu
        novo["ENDEREÇO (REVISAR)"] = ""
        _classificar(numero, cliente, texto_revisar, novo, relatorio)
    else:
        relatorio.sem_endereco += 1


def _detectar_formato(ws, nome_arquivo: str) -> str:
    cabecalho = [_norm(ws.cell(row=1, column=j).value) for j in range(1, ws.max_column + 1)]
    if cabecalho == bd.CLIENTES_COLUNAS:
        raise ErroMigracao(f"A aba CLIENTES de '{nome_arquivo}' já está no formato atual - nada a migrar.")
    for formato, colunas in _COLUNAS_POR_FORMATO.items():
        if cabecalho == colunas:
            return formato
    raise ErroMigracao(
        f"A aba CLIENTES de '{nome_arquivo}' não está em nenhum formato antigo conhecido "
        f"(colunas encontradas: {cabecalho}). Migração recusada para não misturar dados."
    )


def planejar(caminho_xlsx: Path) -> tuple[list[dict], Relatorio]:
    """Le e classifica os enderecos SEM gravar nada (nem backup) - serve pra
    mostrar o que a migracao faria antes de rodar de verdade."""
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=False)
    try:
        ws = wb[bd.ABA_CLIENTES]
        formato = _detectar_formato(ws, caminho_xlsx.name)
        antigos = _ler_registros_antigos(ws, _COLUNAS_POR_FORMATO[formato])
        return _montar_novos_registros(antigos, formato)
    finally:
        wb.close()


def _reescrever_aba(ws, novos: list[dict]) -> None:
    estilo_cabecalho = copy(ws["A1"]._style)
    for j, nome in enumerate(bd.CLIENTES_COLUNAS, start=1):
        celula = ws.cell(row=1, column=j, value=nome)
        celula._style = copy(estilo_cabecalho)
    bd._escrever_linhas_simples(ws, bd.ABA_CLIENTES, bd.CLIENTES_COLUNAS, novos)
    for j, nome in enumerate(bd.CLIENTES_COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = _LARGURAS[nome]


def _valores_da_aba(wb, nome: str) -> list[tuple]:
    return [tuple(linha) for linha in wb[nome].iter_rows(values_only=True)]


def _verificar_resultado(
    caminho_original: Path, caminho_novo: Path, novos: list[dict], relatorio: Relatorio
) -> None:
    """Rele o arquivo novo (ainda temporario) e confere, contra o original,
    que so mudou o que devia mudar. Levanta ErroMigracao se algo destoar."""
    formato = relatorio.formato
    colunas_antigas = _COLUNAS_POR_FORMATO[formato]
    linhas_separadas_agora = {i.linha_planilha for i in relatorio.separados}

    antes = openpyxl.load_workbook(caminho_original, data_only=False)
    depois = openpyxl.load_workbook(caminho_novo, data_only=False)
    try:
        if antes.sheetnames != depois.sheetnames:
            raise ErroMigracao(f"abas mudaram: {antes.sheetnames} -> {depois.sheetnames}")
        for nome in antes.sheetnames:
            if nome != bd.ABA_CLIENTES and _valores_da_aba(antes, nome) != _valores_da_aba(depois, nome):
                raise ErroMigracao(f"a aba {nome} ficou diferente do original")

        ws_novo = depois[bd.ABA_CLIENTES]
        cabecalho = [ws_novo.cell(row=1, column=j).value for j in range(1, ws_novo.max_column + 1)]
        if cabecalho != bd.CLIENTES_COLUNAS:
            raise ErroMigracao(f"cabeçalho novo inesperado: {cabecalho}")

        linhas_novas = [
            dict(zip(bd.CLIENTES_COLUNAS, linha))
            for linha in ws_novo.iter_rows(min_row=2, max_col=len(bd.CLIENTES_COLUNAS), values_only=True)
            if any(v not in (None, "") for v in linha)
        ]
        antigos = _ler_registros_antigos(antes[bd.ABA_CLIENTES], colunas_antigas)
        if len(linhas_novas) != len(antigos):
            raise ErroMigracao(f"{len(antigos)} clientes antes, {len(linhas_novas)} depois")

        colunas_so_novas = [c for c in bd.CLIENTES_COLUNAS if c not in colunas_antigas and c not in _COLUNAS_ENDERECO]
        for (numero, antigo), novo, esperado in zip(antigos, linhas_novas, novos):
            for coluna in _colunas_preservadas(formato):
                if _norm(antigo[coluna]) != _norm(novo[coluna]):
                    raise ErroMigracao(f"linha {numero}, coluna {coluna}: valor mudou na migração")
            # o que foi gravado tem que ser exatamente o que o parser decidiu
            for coluna in _COLUNAS_ENDERECO:
                if _norm(esperado.get(coluna)) != _norm(novo[coluna]):
                    raise ErroMigracao(f"linha {numero}, coluna {coluna}: gravado diferente do planejado")
            # colunas que nao existiam antes (pai/mae/profissao, no formato legado) nascem vazias
            for coluna in colunas_so_novas:
                if _norm(novo[coluna]):
                    raise ErroMigracao(f"linha {numero}: {coluna} deveria estar vazio")

            if formato == FORMATO_LEGADO:
                # nada do endereco original pode sumir: ou foi separado, ou esta inteiro em REVISAR
                if not _norm(antigo["ENDEREÇO"]) and any(_norm(novo[c]) for c in _COLUNAS_ENDERECO):
                    raise ErroMigracao(f"linha {numero}: sem endereço antes, mas com endereço depois")
            elif numero not in linhas_separadas_agora:
                # linha que nao foi re-separada: campos antigos tem que estar intactos
                for coluna in ("CEP", "LOGRADOURO", "NÚMERO", "BAIRRO", "ENDEREÇO (REVISAR)"):
                    if _norm(antigo[coluna]) != _norm(novo[coluna]):
                        raise ErroMigracao(f"linha {numero}, coluna {coluna}: valor mudou na migração")
                cidade_antiga = _norm(antigo["CIDADE"])
                cidade_nova, uf_nova = _norm(novo["CIDADE"]), _norm(novo["UF"])
                reconstruida = f"{cidade_nova}/{uf_nova}" if uf_nova else cidade_nova
                if reconstruida != cidade_antiga:
                    raise ErroMigracao(f"linha {numero}: cidade '{cidade_antiga}' virou '{reconstruida}'")
        # tem que continuar lendo pelo app, sem erro
        bd.ler_clientes(caminho_novo)
    finally:
        antes.close()
        depois.close()


def migrar_planilha(caminho_xlsx: Path, sincronizar: bool = False) -> Relatorio:
    caminho_xlsx = Path(caminho_xlsx)
    if bd.arquivo_esta_bloqueado(caminho_xlsx):
        raise bd.ErroArquivoBloqueado(
            f"O arquivo '{caminho_xlsx.name}' está aberto no Excel. Feche-o e rode a migração de novo."
        )

    novos, relatorio = planejar(caminho_xlsx)  # tambem recusa formato desconhecido / ja migrado

    # 1) backup, conferido byte a byte, ANTES de qualquer escrita
    carimbo = datetime.now().strftime("%Y%m%d-%H%M%S")
    rotulo = _ROTULO_BACKUP[relatorio.formato]
    backup = caminho_xlsx.with_name(f"{caminho_xlsx.stem}.backup-{rotulo}-{carimbo}{caminho_xlsx.suffix}")
    shutil.copy2(caminho_xlsx, backup)
    if _hash_arquivo(backup) != _hash_arquivo(caminho_xlsx):
        backup.unlink(missing_ok=True)
        raise ErroMigracao("O backup não ficou idêntico ao original - migração cancelada, nada foi alterado.")
    relatorio.caminho_backup = backup

    # 2) monta o resultado num arquivo temporario e o confere; o real ainda nao foi tocado
    temporario = caminho_xlsx.with_name(f"{caminho_xlsx.stem}.migrando{caminho_xlsx.suffix}")
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=False)
    try:
        _reescrever_aba(wb[bd.ABA_CLIENTES], novos)
        wb.save(temporario)
    finally:
        wb.close()
    try:
        _verificar_resultado(caminho_xlsx, temporario, novos, relatorio)
        # 3) troca atomica
        os.replace(temporario, caminho_xlsx)
    except PermissionError as exc:
        raise bd.ErroArquivoBloqueado(
            f"Não foi possível substituir '{caminho_xlsx.name}' - ele está aberto em algum programa? "
            f"O backup continua em {backup.name}."
        ) from exc
    finally:
        temporario.unlink(missing_ok=True)

    if sincronizar:
        from core import sheets_sync

        sheets_sync._sincronizar_agora(bd.ABA_CLIENTES, bd.ler_clientes(caminho_xlsx))
    return relatorio
