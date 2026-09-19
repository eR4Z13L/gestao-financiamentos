"""Camada de leitura/escrita do arquivo .xlsx que funciona como banco de dados.

Esta e a UNICA camada que sabe que os dados estao num .xlsx. As camadas de
regra de negocio (core/clientes.py, core/equipamentos.py, core/propostas.py)
so enxergam DataFrames - se um dia o banco virar SQLite/Postgres, so este
arquivo precisa mudar.

Regras importantes:
- O arquivo pode ser aberto manualmente no Excel a qualquer momento.
- Toda escrita rele o arquivo do zero antes de aplicar a mudanca, para nao
  sobrescrever uma edicao feita direto no Excel entre uma tela e outra.
- Toda escrita e feita em arquivo temporario + substituicao atomica, para
  nunca deixar o .xlsx pela metade se algo falhar no meio do caminho.
- Se o Excel estiver com o arquivo aberto, a escrita falha com
  ErroArquivoBloqueado em vez de travar o app - quem chamou decide como
  avisar o usuario (nunca deixamos o erro passar em silencio).
- Toda escrita bem sucedida dispara, em background, uma tentativa de
  sincronizar a aba correspondente com o Google Sheets (core/sheets_sync.py)
  - nunca bloqueia nem falha a escrita local por causa disso.
"""

from __future__ import annotations

import logging
import os
import tempfile
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd

from core import sheets_sync
from core.validators import apenas_digitos

_logger = logging.getLogger(__name__)

ABA_CLIENTES = "CLIENTES"
ABA_EQUIPAMENTOS = "EQUIPAMENTOS"
ABA_PROPOSTAS = "PROPOSTAS"
ABA_VENDEDORES = "VENDEDORES"

# As colunas A-E (DATA CADASTRO, CPF/CNPJ, VENDEDOR, TIPO, CLIENTE) NAO podem
# mudar de lugar: as formulas de VENDEDOR/CLIENTE na aba PROPOSTAS fazem
# VLOOKUP nelas por letra de coluna (CLIENTES!$B:$C e $B:$E) - ver
# _escrever_linhas_propostas.
CLIENTES_COLUNAS = [
    "DATA CADASTRO",
    "CPF/CNPJ",
    "VENDEDOR",
    "TIPO",
    "CLIENTE",
    "NASCIMENTO",
    "CELULAR",
    "CEP",
    "LOGRADOURO",
    "NÚMERO",
    "COMPLEMENTO",
    "BAIRRO",
    "CIDADE",
    "UF",
    # texto do antigo campo unico "ENDEREÇO" de quem a migracao nao conseguiu
    # separar com seguranca (core/endereco.py) - fica so pra revisao manual;
    # apague o conteudo depois de preencher CEP/LOGRADOURO/... daquele cliente
    "ENDEREÇO (REVISAR)",
    "VINCULADO",
    "REDE SOCIAL",
    "EMAIL",
    "NOME DO PAI",
    "NOME DA MÃE",
    "PROFISSÃO",
]

EQUIPAMENTOS_COLUNAS = [
    "FORNECEDOR",
    "EQUIPAMENTO",
    "PARCELAS",
    "VALOR PARCELA (R$)",
    "VALOR LÍQUIDO/REFERÊNCIA (R$)",
    "OBSERVAÇÕES",
]

# Cadastro proprio de vendedores - existe pra nao depender de "quem ja foi
# usado em CLIENTES" (uma aba a mais). SENHA_HASH/SALT sao do login da Fase 2
# (core/auth.py) - a senha em texto puro NUNCA e gravada em lugar nenhum.
VENDEDORES_COLUNAS = ["NOME", "SENHA_HASH", "SALT"]

# Colunas realmente digitadas/gravadas na aba PROPOSTAS.
PROPOSTAS_COLUNAS_EDITAVEIS = [
    "DATA",
    "CPF",
    "VALOR (R$)",
    "MESES",
    "EQUIPAMENTO",
    "BANCO",
    "STATUS",
    "OBSERVAÇÕES",
]

# Ordem completa da aba, incluindo as colunas calculadas (VENDEDOR, CLIENTE,
# TEMPO) que nunca sao digitadas - sao sempre derivadas do CPF e da data.
# O vendedor de uma proposta e sempre o mesmo vendedor do cliente (cadastrado
# uma vez em CLIENTES); repetir esse campo por proposta so criaria risco de
# ficar desatualizado se o cliente trocar de vendedor.
PROPOSTAS_COLUNAS = [
    "DATA",
    "VENDEDOR",
    "CPF",
    "CLIENTE",
    "VALOR (R$)",
    "MESES",
    "EQUIPAMENTO",
    "BANCO",
    "TEMPO",
    "STATUS",
    "OBSERVAÇÕES",
]

# Palavras (em CAIXA ALTA) que contam como "negado" pro calculo de TEMPO e
# pras categorias do dashboard (core/propostas.py reaproveita esta lista, em
# vez de manter a propria copia, pra nunca ficar dessincronizada daqui).
PALAVRAS_STATUS_NEGADO = {"NEGADO", "NEGADA", "REPROVADO", "REPROVADA", "CANCELADO", "CANCELADA"}

# "Efetivado" = aprovada pelo banco E a compra foi de fato concluida - e a etapa
# que vem DEPOIS de "Aprovado" no fluxo (Aprovado -> Efetivado), nao um status
# paralelo. core/propostas.py reaproveita esta lista.
PALAVRAS_STATUS_EFETIVADO = {"EFETIVADO", "EFETIVADA"}

# Status que fazem uma proposta parar de contar "N dias" e virar "Encerrado"
# na coluna TEMPO: os desfechos finais - a compra concluida (Efetivado) ou a
# proposta perdida (negado/reprovado/cancelado). "Aprovado" NAO encerra: a
# proposta segue "em aberto" (contando dias) ate virar Efetivado, ou ser revertida.
STATUS_ENCERRADO = PALAVRAS_STATUS_EFETIVADO | PALAVRAS_STATUS_NEGADO

_COLUNAS_DE_DATA = {
    ABA_CLIENTES: {"DATA CADASTRO", "NASCIMENTO"},
    ABA_PROPOSTAS: {"DATA"},
}
_COLUNAS_DE_MOEDA = {
    ABA_EQUIPAMENTOS: {"VALOR PARCELA (R$)", "VALOR LÍQUIDO/REFERÊNCIA (R$)"},
    ABA_PROPOSTAS: {"VALOR (R$)"},
}

# tudo que nao e data (CELULAR incluido: e sempre texto, nunca numero)
_CLIENTES_COLUNAS_TEXTO = [c for c in CLIENTES_COLUNAS if c not in {"DATA CADASTRO", "NASCIMENTO"}]

# colunas que o Excel deve tratar como TEXTO ao digitar direto na planilha -
# senao um CEP "02378-255" digitado sem o hifen vira o numero 2378255 (perde o
# zero da frente) e "45A" vira erro
_COLUNAS_FORCADAS_A_TEXTO = {
    ABA_CLIENTES: {"CEP", "NÚMERO"},
}
_EQUIPAMENTOS_COLUNAS_TEXTO = ["FORNECEDOR", "EQUIPAMENTO", "OBSERVAÇÕES"]
_EQUIPAMENTOS_COLUNAS_NUMERICAS = ["PARCELAS", "VALOR PARCELA (R$)", "VALOR LÍQUIDO/REFERÊNCIA (R$)"]
_PROPOSTAS_COLUNAS_TEXTO = ["CPF", "EQUIPAMENTO", "BANCO", "STATUS", "OBSERVAÇÕES"]
_PROPOSTAS_COLUNAS_NUMERICAS = ["VALOR (R$)", "MESES"]
_VENDEDORES_COLUNAS_TEXTO = ["NOME", "SENHA_HASH", "SALT"]


class ErroArquivoBloqueado(Exception):
    """O .xlsx esta aberto no Excel (ou outro programa) e nao pode ser salvo agora."""


class ErroPlanilhaDesatualizada(Exception):
    """A aba CLIENTES do .xlsx esta num formato de colunas diferente do que o
    app espera (ex.: planilha de antes da separacao do endereco) - ler ou
    gravar assim trocaria dados de coluna sem ninguem perceber."""


# ---------------------------------------------------------------------------
# Deteccao de bloqueio / leitura basica
# ---------------------------------------------------------------------------

def _caminho_arquivo_bloqueio(caminho_xlsx: Path) -> Path:
    return caminho_xlsx.with_name(f"~${caminho_xlsx.name}")


def arquivo_esta_bloqueado(caminho_xlsx: Path) -> bool:
    """True se tudo indica que o Excel esta com o arquivo aberto agora."""
    return _caminho_arquivo_bloqueio(caminho_xlsx).exists()


def _carregar_planilha(caminho_xlsx: Path):
    return openpyxl.load_workbook(caminho_xlsx, data_only=False)


def _esta_vazio(valor) -> bool:
    """None, "" ou NaN (celula vazia as vezes vira float('nan') em vez de None,
    dependendo de como o pandas infere o tipo da coluna)."""
    if valor is None or valor == "":
        return True
    if isinstance(valor, float):
        return pd.isna(valor)
    return False


def _normalizar_texto(valor) -> str:
    if _esta_vazio(valor):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        # celula numerica (ex: CPF digitado sem formatacao, Excel guarda como
        # numero) vira "11144477735.0" sem isso - corrompendo o texto com um
        # ".0" no final em vez do numero inteiro esperado.
        valor = int(valor)
    return str(valor).strip()


def _eh_formula(valor) -> bool:
    return isinstance(valor, str) and valor.startswith("=")


def _avisar_formulas_nao_calculadas(df: pd.DataFrame, nome_aba: str, colunas: list[str]) -> None:
    """openpyxl (data_only=False) le uma celula com formula como o TEXTO da
    formula (ex: "=A2*1,1"), nunca o valor calculado - convertida pra numero/
    data depois, essa celula vira NaN silenciosamente, como se estivesse
    vazia. Aqui so avisamos no log (nao ha como calcular a formula sem abrir
    o arquivo no Excel), pra quem for investigar um numero "faltando" saber
    que a causa e uma formula na planilha, nao um campo realmente em branco.
    """
    for coluna in colunas:
        if coluna not in df.columns:
            continue
        for indice, valor in df[coluna].items():
            if _eh_formula(valor):
                _logger.warning(
                    "%s!%s (linha %d da planilha) contém uma fórmula do Excel (%r) em vez de um valor "
                    "pronto - foi tratada como vazia. Abra o arquivo no Excel e cole como valor.",
                    nome_aba, coluna, indice + 2, valor,
                )


def _linhas_da_aba(ws, n_colunas: int):
    for linha in ws.iter_rows(min_row=2, max_col=n_colunas, values_only=True):
        if all(v is None or v == "" for v in linha):
            continue
        yield linha


def _conferir_cabecalho(ws, nome_aba: str, colunas: list[str], nome_arquivo: str) -> None:
    """A leitura e a escrita das abas sao por POSICAO de coluna (nao pelo
    nome do cabecalho). Se a planilha estiver num formato diferente do que o
    app espera, ler assim misturaria colunas sem ninguem perceber - entao
    recusa, dizendo exatamente o que esta diferente."""
    encontrado = [_normalizar_texto(ws.cell(row=1, column=j).value) for j in range(1, len(colunas) + 1)]
    if encontrado == colunas:
        return
    diferentes = [
        f"coluna {j}: esperava '{esperada}', achou '{achada}'"
        for j, (esperada, achada) in enumerate(zip(colunas, encontrado), start=1)
        if esperada != achada
    ]
    raise ErroPlanilhaDesatualizada(
        f"A aba {nome_aba} de '{nome_arquivo}' está num formato diferente do esperado "
        f"({'; '.join(diferentes[:3])}{' ...' if len(diferentes) > 3 else ''}). "
        "Se for uma planilha de formato antigo (antes do endereço em campos separados, do complemento ou da UF), "
        "rode: venv\\Scripts\\python.exe scripts\\migrar_endereco_clientes.py"
    )


def ler_aba(
    caminho_xlsx: Path, nome_aba: str, colunas: list[str], conferir_cabecalho: bool = False
) -> pd.DataFrame:
    wb = _carregar_planilha(caminho_xlsx)
    try:
        ws = wb[nome_aba]
        if conferir_cabecalho:
            _conferir_cabecalho(ws, nome_aba, colunas, caminho_xlsx.name)
        linhas = list(_linhas_da_aba(ws, len(colunas)))
    finally:
        wb.close()
    return pd.DataFrame(linhas, columns=colunas)


# ---------------------------------------------------------------------------
# Leitura tipada por entidade
# ---------------------------------------------------------------------------

def ler_clientes(caminho_xlsx: Path) -> pd.DataFrame:
    df = ler_aba(caminho_xlsx, ABA_CLIENTES, CLIENTES_COLUNAS, conferir_cabecalho=True)
    for col in _CLIENTES_COLUNAS_TEXTO:
        df[col] = df[col].map(_normalizar_texto)
    return df


def ler_equipamentos(caminho_xlsx: Path) -> pd.DataFrame:
    df = ler_aba(caminho_xlsx, ABA_EQUIPAMENTOS, EQUIPAMENTOS_COLUNAS)
    _avisar_formulas_nao_calculadas(df, ABA_EQUIPAMENTOS, _EQUIPAMENTOS_COLUNAS_NUMERICAS)
    for col in _EQUIPAMENTOS_COLUNAS_TEXTO:
        df[col] = df[col].map(_normalizar_texto)
    for col in _EQUIPAMENTOS_COLUNAS_NUMERICAS:
        # celula vazia deve virar NaN "de verdade" (tipo numerico), nao um
        # None solto num objeto - senao a tela exibe o texto literal "None".
        df[col] = pd.to_numeric(df[col], errors="coerce")
    return df


def _garantir_aba_vendedores(wb) -> bool:
    """Cria a aba VENDEDORES se o arquivo ainda nao tem uma (planilhas de
    antes dessa funcionalidade existir), populando com os nomes ja usados em
    CLIENTES - sem duplicar por diferenca de maiusculas/espacos, mantendo a
    primeira grafia encontrada. Retorna True se criou agora (quem chamou
    precisa salvar o arquivo depois)."""
    if ABA_VENDEDORES in wb.sheetnames:
        return False

    ws = wb.create_sheet(ABA_VENDEDORES)
    ws.append(VENDEDORES_COLUNAS)

    indice_vendedor = CLIENTES_COLUNAS.index("VENDEDOR")
    vistos: set[str] = set()
    nomes: list[str] = []
    for linha in wb[ABA_CLIENTES].iter_rows(min_row=2, values_only=True):
        if linha is None or len(linha) <= indice_vendedor:
            continue
        nome = _normalizar_texto(linha[indice_vendedor])
        chave = nome.upper()
        if not nome or chave in vistos:
            continue
        vistos.add(chave)
        nomes.append(nome)

    for nome in sorted(nomes, key=str.upper):
        ws.append([nome])
    return True


def ler_vendedores(caminho_xlsx: Path) -> pd.DataFrame:
    wb = _carregar_planilha(caminho_xlsx)
    try:
        migrou = _garantir_aba_vendedores(wb)
        if migrou:
            _salvar_planilha(wb, caminho_xlsx)
        linhas = list(_linhas_da_aba(wb[ABA_VENDEDORES], len(VENDEDORES_COLUNAS)))
    finally:
        wb.close()
    df = pd.DataFrame(linhas, columns=VENDEDORES_COLUNAS)
    for col in _VENDEDORES_COLUNAS_TEXTO:
        df[col] = df[col].map(_normalizar_texto)
    return df


def _calcular_tempo(valor_data, valor_status) -> str:
    if not isinstance(valor_data, (datetime, date)) or pd.isna(valor_data):
        return ""
    status = _normalizar_texto(valor_status).upper()
    if status in STATUS_ENCERRADO:
        return "Encerrado"
    dia = valor_data.date() if isinstance(valor_data, datetime) else valor_data
    dias = (date.today() - dia).days
    return f"{dias} dias"


def ler_propostas(caminho_xlsx: Path, df_clientes: pd.DataFrame | None = None) -> pd.DataFrame:
    """Le a aba PROPOSTAS e recalcula VENDEDOR/CLIENTE/TEMPO a partir do CPF.

    Os valores calculados NUNCA vem do que esta gravado na planilha (que pode
    estar desatualizado se o arquivo nao foi reaberto no Excel) - sao sempre
    recalculados aqui, do mesmo jeito que as formulas de Excel fariam. Isso
    garante que o dashboard e a ficha do cliente sempre mostrem o vendedor e
    o nome corretos mesmo que o cliente tenha sido editado depois da proposta.
    """
    # A aba tem 11 colunas na ordem de PROPOSTAS_COLUNAS (com VENDEDOR/CLIENTE/
    # TEMPO intercalados como formulas). Le tudo nessa ordem e depois descarta
    # essas 3 colunas calculadas - o texto da formula, se sobrar algo, nunca e
    # usado.
    completo = ler_aba(caminho_xlsx, ABA_PROPOSTAS, PROPOSTAS_COLUNAS)
    df = completo[PROPOSTAS_COLUNAS_EDITAVEIS].copy()
    _avisar_formulas_nao_calculadas(df, ABA_PROPOSTAS, [*_PROPOSTAS_COLUNAS_NUMERICAS, "DATA"])
    for col in _PROPOSTAS_COLUNAS_TEXTO:
        df[col] = df[col].map(_normalizar_texto)
    for col in _PROPOSTAS_COLUNAS_NUMERICAS:
        # celula vazia deve virar NaN "de verdade" (tipo numerico), nao um
        # None solto num objeto - senao a tela exibe o texto literal "None".
        df[col] = pd.to_numeric(df[col], errors="coerce")

    if df_clientes is None:
        df_clientes = ler_clientes(caminho_xlsx)
    # cruza por CPF/CNPJ SEM pontuacao dos dois lados - comparar a string
    # exata deixaria VENDEDOR/CLIENTE em branco sempre que a pontuacao
    # divergisse entre as duas abas (ex: "123.456.789-00" vs "12345678900"),
    # mesmo sendo o mesmo cliente.
    referencia = (
        df_clientes.assign(_CHAVE=df_clientes["CPF/CNPJ"].map(apenas_digitos))
        .drop_duplicates("_CHAVE")
        .set_index("_CHAVE")
    )
    chave_proposta = df["CPF"].map(apenas_digitos)
    df["VENDEDOR"] = chave_proposta.map(referencia["VENDEDOR"]).fillna("")
    df["CLIENTE"] = chave_proposta.map(referencia["CLIENTE"]).fillna("")
    df["TEMPO"] = [_calcular_tempo(d, s) for d, s in zip(df["DATA"], df["STATUS"])]
    return df[PROPOSTAS_COLUNAS]


# ---------------------------------------------------------------------------
# Escrita segura (bloqueio + salvamento atomico)
# ---------------------------------------------------------------------------

def _salvar_planilha(wb, caminho_xlsx: Path) -> None:
    if arquivo_esta_bloqueado(caminho_xlsx):
        raise ErroArquivoBloqueado(
            f"O arquivo '{caminho_xlsx.name}' esta aberto no Excel. "
            "Feche-o e tente salvar novamente."
        )
    tmp_fd, tmp_nome = tempfile.mkstemp(suffix=".xlsx", dir=str(caminho_xlsx.parent))
    os.close(tmp_fd)
    caminho_temporario = Path(tmp_nome)
    try:
        wb.save(caminho_temporario)
        os.replace(caminho_temporario, caminho_xlsx)
    except PermissionError as exc:
        caminho_temporario.unlink(missing_ok=True)
        raise ErroArquivoBloqueado(
            f"Nao foi possivel salvar '{caminho_xlsx.name}'. "
            "Verifique se ele nao esta aberto no Excel e tente novamente."
        ) from exc


def _aplicar_formato(cell, nome_aba: str, nome_coluna: str) -> None:
    if nome_coluna in _COLUNAS_DE_DATA.get(nome_aba, ()):
        cell.number_format = "dd/mm/yyyy"
    elif nome_coluna in _COLUNAS_DE_MOEDA.get(nome_aba, ()):
        cell.number_format = '"R$ "#,##0.00'
    elif nome_coluna in _COLUNAS_FORCADAS_A_TEXTO.get(nome_aba, ()):
        cell.number_format = "@"


def _limpar_valor(valor):
    """None/""/NaN viram None (celula vazia de verdade); o resto passa direto."""
    return None if _esta_vazio(valor) else valor


def _escrever_linhas_simples(ws, nome_aba: str, colunas: list[str], registros: list[dict]) -> None:
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)
    for i, registro in enumerate(registros, start=2):
        for j, nome_coluna in enumerate(colunas, start=1):
            cell = ws.cell(row=i, column=j, value=_limpar_valor(registro.get(nome_coluna)))
            _aplicar_formato(cell, nome_aba, nome_coluna)


def _escrever_linhas_propostas(ws, registros: list[dict]) -> None:
    """Regrava a aba PROPOSTAS, incluindo as formulas de VENDEDOR/CLIENTE/TEMPO
    (identicas as que ja existiam na planilha original), para que o arquivo
    continue funcionando normalmente se for aberto e editado direto no Excel.
    """
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)
    for i, registro in enumerate(registros, start=2):
        cell_data = ws.cell(row=i, column=1, value=_limpar_valor(registro.get("DATA")))
        cell_data.number_format = "dd/mm/yyyy"
        # VENDEDOR e CLIENTE sao sempre formulas (VLOOKUP pelo CPF na aba
        # CLIENTES) - nunca gravamos um valor fixo aqui, senao a planilha
        # ficaria desatualizada se o cliente mudasse de vendedor ou de nome.
        ws.cell(row=i, column=2, value=f'=IFERROR(VLOOKUP(C{i},CLIENTES!$B:$C,2,FALSE()),"")')
        ws.cell(row=i, column=3, value=registro.get("CPF"))
        ws.cell(row=i, column=4, value=f'=IFERROR(VLOOKUP(C{i},CLIENTES!$B:$E,4,FALSE()),"")')
        cell_valor = ws.cell(row=i, column=5, value=_limpar_valor(registro.get("VALOR (R$)")))
        cell_valor.number_format = '"R$ "#,##0.00'
        ws.cell(row=i, column=6, value=_limpar_valor(registro.get("MESES")))
        ws.cell(row=i, column=7, value=registro.get("EQUIPAMENTO"))
        ws.cell(row=i, column=8, value=registro.get("BANCO"))
        ws.cell(row=i, column=9, value=_formula_tempo(i))
        ws.cell(row=i, column=10, value=registro.get("STATUS"))
        ws.cell(row=i, column=11, value=registro.get("OBSERVAÇÕES"))


def _formula_tempo(linha: int) -> str:
    """Formula da coluna TEMPO (I) da linha `linha` da aba PROPOSTAS: enquanto
    o status nao for um desfecho final (STATUS_ENCERRADO), mostra "N dias" desde
    o envio; depois disso mostra "Encerrado" (a proposta parou de "correr").
    As condicoes saem do MESMO conjunto usado pelo app (_calcular_tempo), pra
    nunca ficar dessincronizada se um status "final" for adicionado/removido."""
    condicoes_encerrado = ",".join(f'UPPER(J{linha})="{palavra}"' for palavra in sorted(STATUS_ENCERRADO))
    return (
        f'=IF(A{linha}="","",IF(OR({condicoes_encerrado}),'
        f'"Encerrado",TODAY()-A{linha}&" dias"))'
    )


def atualizar_formulas_tempo(caminho_xlsx: Path, gravar: bool = True) -> int:
    """Regrava SO a formula da coluna TEMPO de cada proposta, com a regra de
    "Encerrado" atual. Serve pra planilha que foi gravada com uma regra antiga
    (o app so reescreve as formulas quando alguem grava uma proposta - ate la,
    aberta direto no Excel, ela mostraria a regra velha). Nenhuma outra celula
    e tocada. Devolve quantas formulas mudaram (0 = ja estava em dia); com
    gravar=False so conta, sem salvar nada."""
    wb = _carregar_planilha(caminho_xlsx)
    try:
        ws = wb[ABA_PROPOSTAS]
        if _normalizar_texto(ws.cell(row=1, column=9).value) != "TEMPO":
            raise ErroPlanilhaDesatualizada(
                f"A aba {ABA_PROPOSTAS} de '{caminho_xlsx.name}' não tem a coluna TEMPO na posição esperada (I)."
            )
        alteradas = 0
        for linha in range(2, ws.max_row + 1):
            celula = ws.cell(row=linha, column=9)
            # so mexe em quem ja e formula: uma celula digitada a mao fica como esta
            if _eh_formula(celula.value) and celula.value != _formula_tempo(linha):
                celula.value = _formula_tempo(linha)
                alteradas += 1
        if alteradas and gravar:
            _salvar_planilha(wb, caminho_xlsx)
        return alteradas
    finally:
        wb.close()


def escrever_clientes(caminho_xlsx: Path, df: pd.DataFrame) -> None:
    wb = _carregar_planilha(caminho_xlsx)
    try:
        _conferir_cabecalho(wb[ABA_CLIENTES], ABA_CLIENTES, CLIENTES_COLUNAS, caminho_xlsx.name)
        _escrever_linhas_simples(wb[ABA_CLIENTES], ABA_CLIENTES, CLIENTES_COLUNAS, df.to_dict("records"))
        _salvar_planilha(wb, caminho_xlsx)
    finally:
        wb.close()
    sheets_sync.sincronizar_em_background(ABA_CLIENTES, ler_clientes(caminho_xlsx))


def escrever_equipamentos(caminho_xlsx: Path, df: pd.DataFrame) -> None:
    wb = _carregar_planilha(caminho_xlsx)
    try:
        _escrever_linhas_simples(wb[ABA_EQUIPAMENTOS], ABA_EQUIPAMENTOS, EQUIPAMENTOS_COLUNAS, df.to_dict("records"))
        _salvar_planilha(wb, caminho_xlsx)
    finally:
        wb.close()
    sheets_sync.sincronizar_em_background(ABA_EQUIPAMENTOS, ler_equipamentos(caminho_xlsx))


def escrever_vendedores(caminho_xlsx: Path, df: pd.DataFrame) -> None:
    wb = _carregar_planilha(caminho_xlsx)
    try:
        _garantir_aba_vendedores(wb)
        _escrever_linhas_simples(wb[ABA_VENDEDORES], ABA_VENDEDORES, VENDEDORES_COLUNAS, df.to_dict("records"))
        _salvar_planilha(wb, caminho_xlsx)
    finally:
        wb.close()
    sheets_sync.sincronizar_em_background(ABA_VENDEDORES, ler_vendedores(caminho_xlsx))


def escrever_propostas(caminho_xlsx: Path, df: pd.DataFrame) -> None:
    """`df` deve conter apenas as colunas editaveis (PROPOSTAS_COLUNAS_EDITAVEIS)."""
    wb = _carregar_planilha(caminho_xlsx)
    try:
        _escrever_linhas_propostas(wb[ABA_PROPOSTAS], df.to_dict("records"))
        _salvar_planilha(wb, caminho_xlsx)
    finally:
        wb.close()
    # sincroniza a VISAO COMPLETA (com VENDEDOR/CLIENTE/TEMPO ja calculados,
    # nao formulas) - e o formato que a leitura remota (Fase 2) espera, sem
    # precisar reimplementar as formulas do Excel do outro lado.
    sheets_sync.sincronizar_em_background(ABA_PROPOSTAS, ler_propostas(caminho_xlsx))
