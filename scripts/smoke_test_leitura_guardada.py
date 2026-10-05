"""Valida a leitura guardada na memoria de core/data_store.py (o arquivo so e aberto de novo quando muda):
1. Ler varias vezes, varias abas, abre o arquivo uma vez so.
2. Mexer na tabela devolvida nao estraga a proxima leitura.
3. Gravar pelo app, gravar por fora e trocar o arquivo por outro (mesma data e tamanho) fazem reler.
4. Dois arquivos alternados nunca se misturam; ler menos colunas da mesma aba funciona.
5. O cabecalho errado e recusado em toda leitura, nao so na primeira; aba que nao existe da o mesmo erro de antes.
6. Varias leituras ao mesmo tempo (a sincronizacao le em outra thread) dao o mesmo resultado.

Tudo numa planilha FICTICIA em pasta temporaria (scripts/fixture_ficticia.py); o Google fica bloqueado.

Rodar com: venv/Scripts/python.exe scripts/smoke_test_leitura_guardada.py
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import threading
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config

config.SINCRONIZACAO_GOOGLE_ATIVADA = False  # nunca manda dado de teste pra planilha real na nuvem
import openpyxl

import fixture_ficticia as fx
from core import data_store as bd
from core import data_store_sheets as leitura_sheets
from core import sheets_sync


def linha(titulo: str) -> None:
    print(f"\n{'=' * 60}\n{titulo}\n{'=' * 60}")


class ContarAberturas:
    """Conta quantas vezes o data_store abre o arquivo inteiro (openpyxl.load_workbook)."""

    def __init__(self):
        self.vezes = 0
        self._original = bd.openpyxl.load_workbook

    def __enter__(self):
        def _contando(*args, **kwargs):
            self.vezes += 1
            return self._original(*args, **kwargs)

        bd.openpyxl.load_workbook = _contando
        return self

    def __exit__(self, *_exc):
        bd.openpyxl.load_workbook = self._original


def _versao_nova(caminho: Path) -> None:
    """Muda so a data do arquivo (1 s pra frente): para a leitura guardada, e outra versao."""
    info = os.stat(caminho)
    os.utime(caminho, ns=(info.st_atime_ns, info.st_mtime_ns + 1_000_000_000))


def _nome_do_primeiro_cliente(caminho: Path) -> str:
    return bd.ler_clientes(caminho).iloc[0]["CLIENTE"]


def testar_abre_uma_vez(caminho: Path) -> None:
    linha("1) Varias leituras abrem o arquivo uma vez so")
    _versao_nova(caminho)  # criar a fixture ja deixou esta versao guardada
    with ContarAberturas() as aberturas:
        clientes = bd.ler_clientes(caminho)
        for _ in range(5):
            assert bd.ler_clientes(caminho).equals(clientes), "a leitura repetida devolveu outra coisa"
        bd.ler_propostas(caminho)
        bd.ler_vendas(caminho)
        bd.ler_historico(caminho)
        bd.ler_equipamentos(caminho)
    assert aberturas.vezes == 1, f"esperava abrir o arquivo 1 vez, abriu {aberturas.vezes}"
    print("OK: 5 leituras de CLIENTES e mais PROPOSTAS, VENDAS, HISTÓRICO e EQUIPAMENTOS abriram o arquivo 1 vez.")


def testar_tabela_independente(caminho: Path) -> None:
    linha("2) Mexer na tabela devolvida nao estraga a proxima leitura")
    original = bd.ler_clientes(caminho)
    estragada = bd.ler_clientes(caminho)
    estragada.loc[estragada.index[0], "CLIENTE"] = "Nome Trocado Ficticio"
    estragada.drop(estragada.index[1:], inplace=True)
    assert bd.ler_clientes(caminho).equals(original), "a tabela devolvida antes mudou a leitura seguinte"
    print("OK: trocar um nome e apagar linhas da tabela devolvida nao muda a leitura seguinte.")


def testar_gravacoes_fazem_reler(caminho: Path, pasta: Path) -> None:
    linha("3) Gravar pelo app, gravar por fora e trocar o arquivo fazem reler")
    clientes = bd.ler_clientes(caminho)
    clientes.loc[clientes.index[0], "CLIENTE"] = "Cliente Gravado Pelo App"
    bd.escrever_clientes(caminho, clientes)
    assert _nome_do_primeiro_cliente(caminho) == "Cliente Gravado Pelo App", "a gravacao do app nao apareceu na leitura"
    print("OK: o que o app grava aparece na leitura seguinte.")

    wb = openpyxl.load_workbook(caminho)
    wb[bd.ABA_CLIENTES].cell(row=2, column=bd.CLIENTES_COLUNAS.index("CLIENTE") + 1).value = "Cliente Gravado Por Fora"
    wb.save(caminho)  # no mesmo arquivo, como o Excel faria
    assert _nome_do_primeiro_cliente(caminho) == "Cliente Gravado Por Fora", "a gravacao de fora nao apareceu na leitura"
    print("OK: o que e gravado por fora (no mesmo arquivo) aparece na leitura seguinte.")

    # outro arquivo com o MESMO tamanho e a MESMA data no lugar deste: so a identidade do arquivo muda
    wb = openpyxl.load_workbook(caminho)
    wb[bd.ABA_CLIENTES].cell(row=2, column=bd.CLIENTES_COLUNAS.index("CLIENTE") + 1).value = "Cliente Gravado Por Fera"
    outro = pasta / "outro_do_mesmo_tamanho.xlsx"
    wb.save(outro)
    info = os.stat(caminho)
    if os.stat(outro).st_size != info.st_size:
        print("AVISO: o arquivo trocado saiu com outro tamanho; o caso 'mesma data e tamanho' ficou sem conferir.")
    os.utime(outro, ns=(info.st_atime_ns, info.st_mtime_ns))
    os.replace(outro, caminho)
    assert _nome_do_primeiro_cliente(caminho) == "Cliente Gravado Por Fera", (
        "trocar o arquivo por outro de mesma data e tamanho nao fez reler"
    )
    print("OK: trocar o arquivo por outro de mesma data e tamanho tambem faz reler.")


def testar_dois_arquivos(caminho: Path, pasta: Path) -> None:
    linha("4) Dois arquivos alternados nunca se misturam")
    segundo = pasta / "segundo.xlsx"
    shutil.copy(caminho, segundo)
    clientes = bd.ler_clientes(segundo)
    clientes.loc[clientes.index[0], "CLIENTE"] = "Cliente Do Segundo Arquivo"
    bd.escrever_clientes(segundo, clientes)
    for _ in range(2):
        assert _nome_do_primeiro_cliente(caminho) != "Cliente Do Segundo Arquivo", "leu o segundo arquivo no lugar do primeiro"
        assert _nome_do_primeiro_cliente(segundo) == "Cliente Do Segundo Arquivo", "leu o primeiro arquivo no lugar do segundo"
    print("OK: lendo um arquivo e outro alternadamente, cada leitura traz o seu.")

    # pen drive ou pasta de rede: o Windows pode dar a mesma "identidade" (0) a todo arquivo; com data e
    # tamanho iguais tambem, so o caminho separa os dois
    stat_original = bd.os.stat

    def _stat_sem_identidade(caminho_stat, *a, **k):
        info = stat_original(caminho_stat, *a, **k)
        return os.stat_result((info.st_mode, 0, info.st_dev, info.st_nlink, info.st_uid, info.st_gid, 1000,
                               0, 0, 0, 0, 0, 0))  # tamanho 1000, datas e identidade 0 pra todo arquivo

    bd.os.stat = _stat_sem_identidade
    try:
        for _ in range(2):
            assert _nome_do_primeiro_cliente(caminho) != "Cliente Do Segundo Arquivo", "sem identidade, misturou os arquivos"
            assert _nome_do_primeiro_cliente(segundo) == "Cliente Do Segundo Arquivo", "sem identidade, misturou os arquivos"
    finally:
        bd.os.stat = stat_original
    print("OK: mesmo com data, tamanho e identidade iguais (pen drive, rede), o caminho separa os dois arquivos.")

    poucas = bd.ler_aba(caminho, bd.ABA_CLIENTES, bd.CLIENTES_COLUNAS[:3])
    assert list(poucas.columns) == bd.CLIENTES_COLUNAS[:3] and len(poucas) == len(bd.ler_clientes(caminho))
    print("OK: ler so as 3 primeiras colunas da mesma aba (ja guardada inteira) funciona.")


def testar_cabecalho_e_aba_inexistente(pasta: Path) -> None:
    linha("5) Cabecalho errado recusado sempre; aba inexistente da o mesmo erro de antes")
    errada = pasta / "cabecalho_errado.xlsx"
    shutil.copy(fx.EXEMPLO, errada)
    wb = openpyxl.load_workbook(errada)
    wb[bd.ABA_CLIENTES].cell(row=1, column=2).value = "COLUNA TROCADA"
    wb.save(errada)
    for vez in (1, 2):
        try:
            bd.ler_clientes(errada)
        except bd.ErroPlanilhaDesatualizada as erro:
            assert "COLUNA TROCADA" in str(erro).upper()
        else:
            raise AssertionError(f"a leitura {vez} aceitou o cabecalho errado")
    print("OK: o cabecalho errado e recusado na 1a e na 2a leitura.")

    try:
        bd.ler_aba(errada, "ABA QUE NAO EXISTE", ["A"])
    except KeyError as erro:
        assert "ABA QUE NAO EXISTE" in str(erro)
    else:
        raise AssertionError("ler uma aba que nao existe deveria dar KeyError")
    print("OK: ler uma aba que nao existe da KeyError, como antes.")


def testar_leituras_simultaneas(caminho: Path) -> None:
    linha("6) Varias leituras ao mesmo tempo")
    esperado = bd.ler_propostas(caminho)
    _versao_nova(caminho)  # as threads disputam a abertura
    erros: list[str] = []

    def _ler():
        try:
            for _ in range(5):
                if not bd.ler_propostas(caminho).equals(esperado):
                    erros.append("resultado diferente")
        except Exception as exc:  # noqa: BLE001 - qualquer erro numa thread vira falha do teste
            erros.append(repr(exc))

    threads = [threading.Thread(target=_ler) for _ in range(6)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert not erros, f"leituras simultaneas falharam: {erros[:3]}"
    print("OK: 6 threads lendo PROPOSTAS ao mesmo tempo deram o mesmo resultado, sem erro.")

    # uma leitura lenta em outra thread enquanto o arquivo muda: sem a trava, ela guardaria as linhas da
    # versao velha por cima das da nova, e a tela passaria a mostrar o dado antigo
    _versao_nova(caminho)
    linhas_original = bd._linhas_da_aba
    lendo = threading.Event()

    def _linhas_devagar(ws, n_colunas):
        if threading.current_thread().name == "leitura-lenta":
            lendo.set()
            threading.Event().wait(0.5)
        return linhas_original(ws, n_colunas)

    bd._linhas_da_aba = _linhas_devagar
    try:
        lenta = threading.Thread(target=bd.ler_clientes, args=(caminho,), name="leitura-lenta")
        lenta.start()
        assert lendo.wait(10), "a leitura lenta nao comecou"
        wb = openpyxl.load_workbook(caminho)
        wb[bd.ABA_CLIENTES].cell(row=2, column=bd.CLIENTES_COLUNAS.index("CLIENTE") + 1).value = "Cliente Da Versao Nova"
        wb.save(caminho)
        _versao_nova(caminho)
        assert _nome_do_primeiro_cliente(caminho) == "Cliente Da Versao Nova", "a versao nova nao apareceu"
        lenta.join()
        assert _nome_do_primeiro_cliente(caminho) == "Cliente Da Versao Nova", (
            "a leitura lenta da versao velha sobrescreveu a nova"
        )
    finally:
        bd._linhas_da_aba = linhas_original
    print("OK: uma leitura lenta em outra thread, durante uma gravacao, nao traz de volta o dado antigo.")


def main() -> None:
    def _sem_rede(*_a, **_k):
        raise AssertionError("o teste tentou acessar o Google de verdade")

    originais = (leitura_sheets._obter_cliente, sheets_sync._obter_cliente)
    leitura_sheets._obter_cliente = sheets_sync._obter_cliente = _sem_rede
    pasta = Path(tempfile.mkdtemp(prefix="teste_leitura_guardada_"))
    try:
        caminho = fx.criar(pasta)
        testar_abre_uma_vez(caminho)
        testar_tabela_independente(caminho)
        testar_gravacoes_fazem_reler(caminho, pasta)
        testar_dois_arquivos(caminho, pasta)
        testar_cabecalho_e_aba_inexistente(pasta)
        testar_leituras_simultaneas(caminho)
        linha("TUDO OK")
    finally:
        leitura_sheets._obter_cliente, sheets_sync._obter_cliente = originais
        shutil.rmtree(pasta, ignore_errors=True)


if __name__ == "__main__":
    main()
