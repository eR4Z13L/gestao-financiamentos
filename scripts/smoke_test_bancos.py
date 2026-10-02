"""Testa o cadastro de Bancos: a aba BANCOS nascendo numa planilha antiga (core/data_store.py), as regras
(core/bancos.py), a sincronizacao da aba nova (inclusive com uma nuvem de antes dela), o formulario de proposta
com a lista fechada e o cartao "Bancos" da aba Cadastros (Administracao).

Tudo numa planilha 100% ficticia em pasta temporaria (scripts/fixture_ficticia.py); o Google fica desligado e
qualquer acesso a ele derruba o teste.

Rodar com: venv/Scripts/python.exe scripts/smoke_test_bancos.py
"""

from __future__ import annotations

import hashlib
import os
import sys
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import gspread
import pandas as pd
from PySide6.QtWidgets import QApplication, QInputDialog, QMessageBox

import config

config.SINCRONIZACAO_GOOGLE_ATIVADA = False
from ambiente_de_teste import Ambiente, Mensagens
from core import bancos as bancos_mod
from core import dashboard as dash
from core import data_store as bd
from core import data_store_sheets as nuvem
from core import propostas as propostas_mod
from core import sessao as sessao_mod
from core import sheets_sync
from core import sincronizacao
from desktop.theme import TEMA_ESCURO
from desktop.widgets.formulario_proposta import FormularioProposta


def linha(titulo: str) -> None:
    print(f"\n{'=' * 60}\n{titulo}\n{'=' * 60}")


def _hash(caminho: Path) -> str:
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


def _banco_das_propostas() -> list[str]:
    return bd.ler_propostas(bancos_mod.CAMINHO_XLSX)["BANCO"].tolist()


def _definir_bancos_das_propostas(nomes: list[str]) -> None:
    df = bd.ler_propostas(bancos_mod.CAMINHO_XLSX)[bd.PROPOSTAS_COLUNAS_EDITAVEIS]
    assert len(df) >= len(nomes)
    for i, nome in enumerate(nomes):
        df.at[i, "BANCO"] = nome
    bd.escrever_propostas(bancos_mod.CAMINHO_XLSX, df)


def main() -> None:
    app = QApplication.instance() or QApplication(sys.argv)
    amb = Ambiente()
    caminho = amb.caminho
    try:
        with Mensagens() as msgs:
            linha("1) Planilha sem a aba BANCOS: ela nasce com a lista inicial + os bancos das propostas")
            sessao_mod.iniciar(sessao_mod.Sessao(papel=sessao_mod.PAPEL_ADMIN, nome_usuario="Administrador"))
            _definir_bancos_das_propostas(
                ["Banco Exemplo", "BANCO EXEMPLO", "Banco Exemplo", "SANTANDER", "Todos", "", "Outro Banco Exemplo"]
            )
            import openpyxl

            with open(caminho, "rb") as f:
                assert bd.ABA_BANCOS not in openpyxl.load_workbook(f).sheetnames, "a fixture e de antes do cadastro"
            df = bd.ler_bancos(caminho)
            nomes = df["NOME"].tolist()
            esperados = sorted({*config.BANCOS_INICIAIS, "Banco Exemplo", "Outro Banco Exemplo"}, key=str.upper)
            assert nomes == esperados, nomes
            assert (df["ATIVO"] == "Sim").all()
            print(f"OK: {len(nomes)} bancos, sem repetir por maiúsculas; vale a grafia da lista inicial ('Santander') "
                  "ou a mais usada ('Banco Exemplo'); 'Todos' e banco vazio ficam de fora.")
            antes_de_ler = _hash(caminho)
            assert bd.ler_bancos(caminho).equals(df) and _hash(caminho) == antes_de_ler, "ler nunca grava"
            with open(caminho, "rb") as f:
                assert bd.ABA_BANCOS not in openpyxl.load_workbook(f).sheetnames
            print("OK: ler o cadastro nunca grava (a aba ainda não existe no arquivo; a lista é calculada na hora).")

            linha("2) Regras: adicionar, renomear, desativar/reativar, excluir")
            assert bancos_mod.adicionar_banco("  Banco   Novo  ") == "Banco Novo"
            with open(caminho, "rb") as f:
                assert bd.ABA_BANCOS in openpyxl.load_workbook(f).sheetnames, "a 1a gravacao cria a aba"
            assert bancos_mod.listar_bancos()["NOME"].tolist() == sorted([*esperados, "Banco Novo"], key=str.upper)
            for ruim, trecho in (("banco novo", "já está"), ("   ", "Informe"), ("TODOS", "opção fixa")):
                try:
                    bancos_mod.adicionar_banco(ruim)
                    raise AssertionError(f"deveria recusar {ruim!r}")
                except bancos_mod.ErroBanco as exc:
                    assert trecho in str(exc), str(exc)
            print("OK: nome repetido (sem diferenciar maiúsculas), em branco e 'Todos' são recusados.")

            uso = bancos_mod.listar_bancos().set_index("NOME")["PROPOSTAS"]
            assert uso["Banco Exemplo"] >= 3 and uso["Santander"] == 1 and uso["Banco Novo"] == 0
            antes = _banco_das_propostas()
            mudaram = bancos_mod.renomear_banco("banco exemplo", "Banco Exemplo SA")
            depois = _banco_das_propostas()
            do_banco = [b for b in antes if b.upper() == "BANCO EXEMPLO"]
            assert mudaram == len(do_banco) >= 3 and "BANCO EXEMPLO" in do_banco
            assert depois.count("Banco Exemplo SA") == mudaram and "BANCO EXEMPLO" not in depois and "Banco Exemplo" not in depois
            assert [a for a, d in zip(antes, depois) if a != d] == do_banco
            print("OK: renomear troca o cadastro e as propostas (todas as grafias), e nenhuma outra proposta.")
            try:
                bancos_mod.renomear_banco("Banco Exemplo SA", "santander")
                raise AssertionError("renomear para o nome de outro banco deveria ser recusado")
            except bancos_mod.ErroBanco as exc:
                assert "outro banco" in str(exc)
            assert bancos_mod.renomear_banco("Smart", "SMART") == 0 and "SMART" in bancos_mod.nomes_ativos()
            print("OK: não deixa juntar dois bancos renomeando; só mudar maiúsculas vale.")

            bancos_mod.desativar_banco("Banco Exemplo SA")
            assert "Banco Exemplo SA" not in bancos_mod.nomes_ativos()
            assert _banco_das_propostas().count("Banco Exemplo SA") == mudaram, "desativar nao mexe nas propostas"
            bancos_mod.reativar_banco("banco exemplo sa")
            assert "Banco Exemplo SA" in bancos_mod.nomes_ativos()
            print("OK: desativar tira da lista ativa sem mexer nas propostas; reativar volta.")
            try:
                bancos_mod.excluir_banco("Santander")
                raise AssertionError("banco usado em proposta nao pode ser excluido")
            except bancos_mod.ErroBanco as exc:
                assert "Desative" in str(exc)
            bancos_mod.excluir_banco("Banco Novo")
            assert "Banco Novo" not in bancos_mod.listar_bancos()["NOME"].tolist()
            try:
                bancos_mod.excluir_banco("Banco Que Nao Existe")
                raise AssertionError("deveria recusar")
            except bancos_mod.ErroBanco:
                pass
            print("OK: excluir só sem propostas (senão explica: desative); banco inexistente é recusado.")

            sessao_mod.encerrar()
            try:
                bancos_mod.adicionar_banco("Sem Permissao")
                raise AssertionError("sem sessao de Administrador nao pode gravar")
            except sessao_mod.PermissaoNegada:
                pass
            sessao_mod.iniciar(sessao_mod.Sessao(papel=sessao_mod.PAPEL_ADMIN, nome_usuario="Administrador"))
            print("OK: só o Administrador grava no cadastro.")

            linha("3) Sincronização: a aba BANCOS vai para a nuvem; uma nuvem sem ela não apaga a daqui")
            enviados = []
            original_envio = sheets_sync.sincronizar_em_background
            sheets_sync.sincronizar_em_background = lambda aba, df: enviados.append(aba)
            try:
                bancos_mod.adicionar_banco("Banco Sincronizado")
                assert enviados == [bd.ABA_BANCOS], enviados
                enviados.clear()
                sincronizacao.sincronizar_tudo_agora(caminho)
                assert bd.ABA_BANCOS in enviados and len(enviados) == 5, enviados
            finally:
                sheets_sync.sincronizar_em_background = original_envio
            print("OK: gravar o cadastro envia a aba BANCOS; 'Sincronizar agora' envia as 5 abas.")

            leitores = {nome: getattr(nuvem, nome) for nome in
                        ("ler_clientes", "ler_equipamentos", "ler_vendedores", "ler_propostas", "ler_bancos")}
            local = {
                "ler_clientes": bd.ler_clientes(caminho), "ler_equipamentos": bd.ler_equipamentos(caminho),
                "ler_vendedores": bd.ler_vendedores(caminho), "ler_propostas": bd.ler_propostas(caminho),
            }
            try:
                for nome, df_local in local.items():
                    setattr(nuvem, nome, lambda df_local=df_local: df_local.copy())

                def _sem_aba():
                    raise gspread.WorksheetNotFound(bd.ABA_BANCOS)

                nuvem.ler_bancos = _sem_aba
                dados = sincronizacao._ler_dados_da_nuvem()
                assert bd.ABA_BANCOS not in dados and len(dados) == 4
                antes_bancos = bd.ler_bancos(caminho)
                sincronizacao._recusar_se_a_nuvem_parece_vazia_por_engano(dados, caminho)
                sincronizacao._gravar_dados_no_arquivo(caminho, dados)
                assert bd.ler_bancos(caminho).equals(antes_bancos), "nuvem sem BANCOS nao pode mexer no cadastro daqui"
                print("OK: nuvem de antes do cadastro (sem a aba): baixa as 4 abas e o cadastro daqui fica igual.")

                da_nuvem = pd.DataFrame([{"NOME": "Banco Da Nuvem", "ATIVO": "Sim"}, {"NOME": "Santander", "ATIVO": "Não"}])
                nuvem.ler_bancos = lambda: da_nuvem.copy()
                dados = sincronizacao._ler_dados_da_nuvem()
                sincronizacao._gravar_dados_no_arquivo(caminho, dados)
                assert bd.ler_bancos(caminho)[["NOME", "ATIVO"]].values.tolist() == da_nuvem.values.tolist()
                assert bancos_mod.nomes_ativos() == ["Banco Da Nuvem"]
                print("OK: nuvem com a aba BANCOS: o cadastro daqui vira o da nuvem.")

                nuvem.ler_bancos = lambda: pd.DataFrame(columns=bd.BANCOS_COLUNAS)
                try:
                    sincronizacao._recusar_se_a_nuvem_parece_vazia_por_engano(sincronizacao._ler_dados_da_nuvem(), caminho)
                    raise AssertionError("aba BANCOS vazia na nuvem com bancos aqui deveria ser recusada")
                except sincronizacao.ErroNuvem as exc:
                    assert bd.ABA_BANCOS in str(exc)
                print("OK: aba BANCOS vazia na nuvem (envio interrompido) não apaga o cadastro daqui.")
            finally:
                for nome, funcao in leitores.items():
                    setattr(nuvem, nome, funcao)

            linha("4) Dashboard: os bancos 'ainda não tentados' vêm do cadastro (só os ativos)")
            bancos_mod.reativar_banco("Santander")
            bancos_mod.adicionar_banco("Banco Inativo")
            bancos_mod.desativar_banco("Banco Inativo")
            propostas = bd.ler_propostas(caminho)
            propostas["STATUS"] = propostas_mod.STATUS_NEGADO
            sugeridos = {b for c in dash.para_reenviar(propostas) for b in c.bancos_nao_tentados}
            assert "Banco Da Nuvem" in sugeridos and "Banco Inativo" not in sugeridos, sugeridos
            print("OK: banco ativo do cadastro entra nas sugestões; desativado não.")

            linha("5) Formulário de proposta: lista fechada (ativos + 'Todos')")
            nomes_ativos = bancos_mod.nomes_ativos()
            novo = FormularioProposta(cpf=None)
            itens = [novo._banco.itemText(i) for i in range(novo._banco.count())]
            assert itens == [*nomes_ativos, bd.BANCO_TODOS], itens
            assert not novo._banco.isEditable() and novo._banco.currentText() == "", "nada escolhido numa proposta nova"
            print(f"OK: {len(itens)} opções (os ativos + 'Todos'), sem digitação livre; começa em branco.")

            proposta = bd.ler_propostas(caminho).iloc[0].to_dict()
            proposta["BANCO"] = "SANTANDER"
            f = FormularioProposta(cpf=proposta["CPF"], proposta=proposta, indice=0)
            assert f._banco.currentText() == "SANTANDER" and f._banco.count() == len(itens) + 1, "mostra o gravado, como esta"
            assert not f.tem_alteracoes()
            proposta["BANCO"] = "Banco Inativo"
            f = FormularioProposta(cpf=proposta["CPF"], proposta=proposta, indice=0)
            assert f._banco.currentText() == "Banco Inativo" and f._banco.count() == len(itens) + 1
            assert not f.tem_alteracoes(), "abrir uma proposta com banco fora da lista nao conta como alteracao"
            print("OK: a proposta mostra o banco exatamente como gravado ('SANTANDER', 'Banco Inativo'), sem trocar nada.")

            linha("6) Administração > Cadastros > Bancos")
            sessao_mod.encerrar()
            janela = amb.nova_janela("admin", TEMA_ESCURO)
            adm = janela._tela_administracao
            abas = adm.findChild(__import__("PySide6.QtWidgets", fromlist=["QTabWidget"]).QTabWidget)
            assert [abas.tabText(i) for i in range(abas.count())] == ["Usuários", "Sincronização e backup", "Cadastros"]
            assert not hasattr(adm, "_botao_mesclar"), "o 'Mesclar grafias' saiu"
            cad = adm._cadastro_bancos
            modelo = cad._modelo
            nomes_na_tabela = [modelo.indice_real(i) for i in range(modelo.rowCount())]
            assert nomes_na_tabela == bancos_mod.listar_bancos()["NOME"].tolist()
            print(f"OK: aba 'Cadastros' com {len(nomes_na_tabela)} bancos; o 'Mesclar' não existe mais.")

            avisos = []
            adm.dados_atualizados.connect(lambda: avisos.append(1))
            original_texto = QInputDialog.getText
            try:
                QInputDialog.getText = staticmethod(lambda *a, **k: ("Banco Pela Tela", True))
                msgs.limpar()
                cad._botao_novo.click()
                assert "Banco Pela Tela" in bancos_mod.nomes_ativos() and msgs.ultima()[1] == "Banco cadastrado" and avisos
                QInputDialog.getText = staticmethod(lambda *a, **k: ("banco pela tela", True))
                cad._botao_novo.click()
                assert msgs.ultima()[1] == "Não foi possível cadastrar" and "já está" in msgs.ultima()[2]
                print("OK: '+ Novo Banco' cadastra e avisa as outras telas; repetido mostra o motivo.")

                cad._tabela.selectRow([cad._modelo.indice_real(i) for i in range(cad._modelo.rowCount())].index("Santander"))
                assert cad._botao_ativo.text() == "Desativar" and "Usado em" in cad._botao_excluir.toolTip()
                msgs.limpar()
                cad._botao_excluir.click()
                assert msgs.ultima()[1] == "Não foi possível excluir" and "Santander" in bancos_mod.nomes_ativos()
                msgs.resposta_pergunta = QMessageBox.StandardButton.No
                cad._botao_ativo.click()
                assert "Santander" in bancos_mod.nomes_ativos(), "respondendo Não, nada muda"
                msgs.resposta_pergunta = QMessageBox.StandardButton.Yes
                cad._botao_ativo.click()
                assert "Santander" not in bancos_mod.nomes_ativos() and msgs.ultima()[1] == "Banco desativado"
                cad._tabela.selectRow([cad._modelo.indice_real(i) for i in range(cad._modelo.rowCount())].index("Santander"))
                assert cad._botao_ativo.text() == "Reativar"
                cad._botao_ativo.click()
                assert "Santander" in bancos_mod.nomes_ativos()
                print("OK: excluir banco com propostas explica; desativar pergunta antes; reativar volta.")

                bancos_mod.adicionar_banco("Banco Exemplo SA")  # o passo 3 trocou o cadastro pelo da nuvem
                cad.recarregar()
                cad._tabela.selectRow([cad._modelo.indice_real(i) for i in range(cad._modelo.rowCount())].index("Banco Exemplo SA"))
                QInputDialog.getText = staticmethod(lambda *a, **k: ("Banco Exemplo Renomeado", True))
                cad._botao_renomear.click()
                assert msgs.ultima()[1] == "Banco renomeado" and f"{mudaram} proposta(s)" in msgs.ultima()[2]
                assert _banco_das_propostas().count("Banco Exemplo Renomeado") == mudaram
                print("OK: renomear pela tela diz quantas propostas mudaram.")
            finally:
                QInputDialog.getText = original_texto
        linha("TUDO OK")
    finally:
        amb.encerrar()


if __name__ == "__main__":
    main()
