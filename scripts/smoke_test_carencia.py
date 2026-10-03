"""Testa a carencia (dias ate a 1a parcela) e o valor da parcela na proposta: planilha antiga sem as colunas, as
regras (opcionais, mas validas quando preenchidas), o formulario, "Mandar a outro banco" (a carencia vem junto, a
parcela nao), a linha do banco no card da venda e a leitura da nuvem.

Planilha 100% ficticia em pasta temporaria, preferencias isoladas e o Google desligado (qualquer acesso derruba).

Rodar com: venv/Scripts/python.exe scripts/smoke_test_carencia.py
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl
import pandas as pd
from PySide6.QtWidgets import QApplication, QDialog

import config

config.SINCRONIZACAO_GOOGLE_ATIVADA = False
import fixture_ficticia as fx
from ambiente_de_teste import Mensagens
from core import bancos as bancos_mod
from core import data_store as bd
from core import data_store_sheets as nuvem
from core import propostas as propostas_mod
from core import sessao as sessao_mod
from core import sheets_sync
from core import vendas as v
from desktop.dialogs.pedido_venda_dialog import PedidoVendaDialog
from desktop.screens import vendas_screen as tela_mod
from desktop.theme import TEMA_ESCURO, build_stylesheet
from desktop.widgets.formulario_proposta import FormularioProposta

ADMIN = sessao_mod.Sessao(papel=sessao_mod.PAPEL_ADMIN, nome_usuario="Administrador")
CARENCIA, PARCELA = "CARÊNCIA (DIAS)", "PARCELA (R$)"


def linha(titulo: str) -> None:
    print(f"\n{'=' * 60}\n{titulo}\n{'=' * 60}")


def _recusa(trecho: str, campos: dict) -> None:
    try:
        propostas_mod._validar_campos(dict(campos))
    except propostas_mod.ErroProposta as exc:
        assert trecho in str(exc), (trecho, str(exc))
        return
    raise AssertionError(f"deveria recusar: {campos}")


def main() -> None:
    def _sem_rede(*_a, **_k):
        raise AssertionError("o teste tentou acessar o Google de verdade")

    rede_original = (nuvem._obter_cliente, sheets_sync._obter_cliente, nuvem._ler_aba_bruta)
    nuvem._obter_cliente = sheets_sync._obter_cliente = _sem_rede
    pasta = Path(tempfile.mkdtemp(prefix="_smoke_carencia_"))
    registro_antes = fx.instantaneo_do_registro()
    fx.isolar_preferencias(pasta)
    caminho = fx.criar(pasta)
    fx.apontar_modulos_para(caminho)
    sessao_mod.iniciar(ADMIN)
    app = QApplication.instance() or QApplication(sys.argv)
    app.setStyleSheet(build_stylesheet(TEMA_ESCURO))
    try:
        with Mensagens() as msgs:
            b1, b2 = bancos_mod.nomes_ativos()[:2]

            linha("1) Planilha de antes: lê as duas colunas em branco; a 1ª gravação põe o cabeçalho no fim")
            with open(caminho, "rb") as f:  # a planilha como era antes: sem as duas colunas do fim
                wb = openpyxl.load_workbook(f)
            wb[bd.ABA_PROPOSTAS].delete_cols(14, 2)
            wb.save(caminho)
            antes = bd.ler_propostas(caminho)
            assert antes[CARENCIA].isna().all() and antes[PARCELA].isna().all()
            id_venda, id_p = v.criar_venda(fx.CPF_MARIA, ["Cadeira Exemplo"],
                                           {"BANCO": b1, "VALOR (R$)": 80000, "MESES": 48, CARENCIA: "45", PARCELA: "2.100,50"})
            with open(caminho, "rb") as f:
                ws = openpyxl.load_workbook(f)[bd.ABA_PROPOSTAS]
                assert [ws.cell(row=1, column=c).value for c in (14, 15)] == [CARENCIA, PARCELA]
                assert str(ws.cell(row=2, column=2).value).startswith("=IFERROR(VLOOKUP("), "as fórmulas não se deslocam"
            nova = bd.ler_propostas(caminho).loc[v.indice_da_proposta(id_p)]
            assert nova[CARENCIA] == 45 and nova[PARCELA] == 2100.5, (nova[CARENCIA], nova[PARCELA])
            with open(caminho, "rb") as f:  # alguem digitou a carencia no Excel como TEXTO: o app le como numero
                wb = openpyxl.load_workbook(f)
            wb[bd.ABA_PROPOSTAS].cell(row=v.indice_da_proposta(id_p) + 2, column=14, value="60")
            wb.save(caminho)
            lidas = bd.ler_propostas(caminho)
            assert lidas[CARENCIA].dtype.kind == "f" and lidas[PARCELA].dtype.kind == "f", "as duas são números"
            assert lidas.loc[v.indice_da_proposta(id_p), CARENCIA] == 60
            v.mudar_status_proposta(id_p, propostas_mod.STATUS_EM_ANALISE)  # regrava: volta a ser 45 abaixo
            propostas_mod.atualizar_proposta(v.indice_da_proposta(id_p), {CARENCIA: 45})
            print("OK: sem as colunas lê em branco; gravar põe 'CARÊNCIA (DIAS)' e 'PARCELA (R$)' nas colunas 14/15; '2.100,50' vira 2100.5.")

            linha("2) Regras: opcionais, mas válidas quando preenchidas")
            base = {"CPF": fx.CPF_MARIA, "EQUIPAMENTO": "X", "VALOR (R$)": 1000}
            for ruim in ("-1", "366", "30.5", "abc"):
                _recusa("Carência inválida", {**base, CARENCIA: ruim})
            for ruim in (0, "-10", "xyz"):
                _recusa("parcela inválido", {**base, PARCELA: ruim})
            campos = {**base, CARENCIA: "0", PARCELA: ""}
            propostas_mod._validar_campos(campos)
            assert campos[CARENCIA] == 0, "0 dias é uma carência válida"
            propostas_mod._validar_campos({**base, CARENCIA: None, PARCELA: float("nan")})
            print("OK: carência inteira de 0 a 365 dias; parcela maior que zero; as duas podem ficar em branco.")

            linha("3) Formulário: as duas na primeira linha, gravam e voltam na leitura (com copiar)")
            f = FormularioProposta(fx.CPF_JOAO, "JOÃO")
            itens = [f._carencia.itemText(i) for i in range(f._carencia.count())]
            assert itens == ["30", "45", "60", "90"], itens
            assert f._carencia.currentText() == "" and f._parcela.value() == 0, "proposta nova começa em branco"
            f._valor.setValue(50000)
            f._meses.setValue(36)
            f._equipamento.setCurrentText("Mesa Exemplo")
            f._banco.setCurrentText(b1)
            f._carencia.setCurrentText("120")  # um prazo fora da lista: digitado
            f._parcela.setValue(1890.75)
            msgs.limpar()
            f._salvar()
            msgs.exigir_vazio("proposta nova com carência e parcela")
            gravada = bd.ler_propostas(caminho).loc[f.indice_gravado]
            assert gravada[CARENCIA] == 120 and gravada[PARCELA] == 1890.75
            leitura = FormularioProposta(fx.CPF_JOAO, "JOÃO", proposta=gravada.to_dict(), indice=f.indice_gravado)
            assert leitura._carencia.currentText() == "120" and leitura._parcela.value() == 1890.75
            assert leitura._carencia.travado() and leitura._parcela.isReadOnly(), "na leitura, travados como os outros"
            assert leitura._carencia in leitura._campos_editaveis() and leitura._parcela in leitura._campos_editaveis()
            leitura._habilitar_edicao()
            leitura._carencia.setCurrentText("")
            leitura._parcela.setValue(0)
            leitura._salvar()
            limpa = bd.ler_propostas(caminho).loc[f.indice_gravado]
            assert pd.isna(limpa[CARENCIA]) and pd.isna(limpa[PARCELA]), "apagar volta a ficar em branco"
            print("OK: 30/45/60/90 pra escolher ou digitar outro; grava, reabre preenchido e travado; apagar deixa em branco.")

            linha("4) Mandar a outro banco: a carência vem junto, a parcela não")
            dup = propostas_mod.dados_para_duplicar(bd.ler_propostas(caminho).loc[v.indice_da_proposta(id_p)].to_dict())
            assert dup[CARENCIA] == 45 and dup[PARCELA] is None
            tela = tela_mod.VendasScreen()
            resumo = next(r for r in v.listar_vendas() if r.id == id_venda)
            d = PedidoVendaDialog(tela, venda=resumo)
            assert d._carencia.currentText() == "45" and d._parcela.value() == 0
            d._banco.setCurrentText(b2)
            d._parcela.setValue(2050)
            d._gravar()
            assert d.result() == QDialog.DialogCode.Accepted
            segunda = bd.ler_propostas(caminho).loc[v.indice_da_proposta(d.id_proposta)]
            assert segunda[CARENCIA] == 45 and segunda[PARCELA] == 2050
            print("OK: o pedido ao 2º banco já vem com 45 dias; a parcela é a dele (cada banco calcula a sua).")

            linha("5) Card da venda: as condições embaixo do nome do banco")
            tela.carregar()
            card = next(c for c in tela.cards if c.venda.id == id_venda)
            textos = [l["condicoes"].text() for l in card.linhas_de_banco]
            assert textos == ["R$ 80.000 · 48x · 1ª em 45 dias · parcela R$ 2.100,50",
                              "R$ 80.000 · 48x · 1ª em 45 dias · parcela R$ 2.050,00"], textos
            print(f"OK: '{textos[0]}'.")

            linha("6) Nuvem: as duas colunas voltam como número")
            bruto = bd.ler_propostas(caminho).astype(object)
            bruto[CARENCIA] = bruto[CARENCIA].map(lambda x: "" if pd.isna(x) else str(int(x)))
            bruto[PARCELA] = bruto[PARCELA].map(lambda x: "" if pd.isna(x) else str(x))
            nuvem._ler_aba_bruta = lambda _aba: bruto.copy()
            lida = nuvem.ler_propostas()
            assert lida[CARENCIA].dtype.kind == "f" and lida[PARCELA].dtype.kind == "f"
            assert sorted(lida[CARENCIA].dropna()) == [45, 45]
            sem = bruto.drop(columns=[CARENCIA, PARCELA])  # nuvem de antes, sem as colunas
            nuvem._ler_aba_bruta = lambda _aba: sem.copy()
            assert nuvem.ler_propostas()[CARENCIA].isna().all()
            print("OK: da nuvem a carência e a parcela chegam como número; nuvem antiga, sem as colunas, lê em branco.")
        linha("TUDO OK")
    finally:
        nuvem._obter_cliente, sheets_sync._obter_cliente, nuvem._ler_aba_bruta = rede_original
        fx.restaurar_modulos()
        sessao_mod.encerrar()
        assert fx.instantaneo_do_registro() == registro_antes, "o teste mexeu nas preferências reais"


if __name__ == "__main__":
    main()
