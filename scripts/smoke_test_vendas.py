"""Testa a base das vendas (core/vendas.py e as abas VENDAS/HISTÓRICO de core/data_store.py): criar venda com um ou
mais equipamentos, mandar a outro banco, as regras de status da proposta e da venda, o historico, os alertas de 7
dias, planilha antiga sem as colunas novas e a sincronizacao das abas novas (inclusive com uma nuvem sem elas).

Tudo numa planilha 100% ficticia em pasta temporaria (scripts/fixture_ficticia.py); o Google fica desligado.

Rodar com: venv/Scripts/python.exe scripts/smoke_test_vendas.py
"""

from __future__ import annotations

import os
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import gspread
import openpyxl
import pandas as pd

import config

config.SINCRONIZACAO_GOOGLE_ATIVADA = False
import fixture_ficticia as fx
from core import data_store as bd
from core import data_store_sheets as nuvem
from core import propostas as propostas_mod
from core import sessao as sessao_mod
from core import sheets_sync
from core import sincronizacao
from core import vendas as v

ADMIN = sessao_mod.Sessao(papel=sessao_mod.PAPEL_ADMIN, nome_usuario="Administrador")


def linha(titulo: str) -> None:
    print(f"\n{'=' * 60}\n{titulo}\n{'=' * 60}")


def _recusa(trecho: str, funcao, *args, **kwargs) -> None:
    try:
        funcao(*args, **kwargs)
    except v.ErroVenda as exc:
        assert trecho in str(exc), (trecho, str(exc))
        return
    raise AssertionError(f"deveria recusar ({trecho})")


def _pedido(banco: str, valor: float = 50000, **extra) -> dict:
    return {"BANCO": banco, "VALOR (R$)": valor, "MESES": 36, **extra}


def main() -> None:
    def _sem_rede(*_a, **_k):
        raise AssertionError("o teste tentou acessar o Google de verdade")

    rede_original = (nuvem._obter_cliente, sheets_sync._obter_cliente)
    nuvem._obter_cliente = sheets_sync._obter_cliente = _sem_rede
    pasta = Path(tempfile.mkdtemp(prefix="_smoke_vendas_"))
    caminho = fx.criar(pasta)
    fx.apontar_modulos_para(caminho)
    sessao_mod.iniciar(ADMIN)
    try:
        linha("1) Planilha de antes das vendas: lê sem as colunas e abas novas, e editar não estraga nada")
        antes = bd.ler_propostas(caminho)
        assert (antes["ID_PROPOSTA"] == "").all() and (antes["ID_VENDA"] == "").all()
        assert bd.ler_vendas(caminho).empty and bd.ler_historico(caminho).empty
        with open(caminho, "rb") as f:
            assert {bd.ABA_VENDAS, bd.ABA_HISTORICO}.isdisjoint(openpyxl.load_workbook(f).sheetnames), "ler nunca cria aba"
        print("OK: planilha antiga lê normal (IDs em branco, sem abas novas) e ler não grava nada.")

        linha("2) Criar venda: venda + 1ª proposta + histórico, numa gravação só")
        id_venda, id_p1 = v.criar_venda(fx.CPF_MARIA, ["Cadeira Exemplo", "  laser exemplo ", "CADEIRA EXEMPLO"], _pedido("Banco Exemplo"))
        vendas = bd.ler_vendas(caminho)
        assert len(vendas) == 1 and id_venda.startswith("V-") and id_p1.startswith("P-")
        venda = vendas.iloc[0]
        assert venda["EQUIPAMENTOS"] == "Cadeira Exemplo + laser exemplo", "repetido some, espaços saem"
        assert venda["STATUS"] == v.VENDA_AGUARDANDO and venda["BANCO_ESCOLHIDO"] == ""
        p1 = bd.ler_propostas(caminho).set_index("ID_PROPOSTA").loc[id_p1]
        assert p1["ID_VENDA"] == id_venda and p1["EQUIPAMENTO"] == venda["EQUIPAMENTOS"], "sem equipamento: cobre a venda toda"
        assert p1["STATUS"] == propostas_mod.STATUS_EM_ANALISE
        hist = bd.ler_historico(caminho)
        assert list(hist["ID"]) == [id_venda, id_p1] and list(hist["PARA"]) == [v.VENDA_AGUARDANDO, propostas_mod.STATUS_EM_ANALISE]
        assert hist["COMPUTADOR"].iloc[0] == sheets_sync.nome_desta_maquina() and hist["QUANDO"].notna().all()
        with open(caminho, "rb") as f:
            ws = openpyxl.load_workbook(f)[bd.ABA_PROPOSTAS]
            assert [ws.cell(row=1, column=j).value for j in (12, 13)] == ["ID_PROPOSTA", "ID_VENDA"]
        print("OK: venda com 2 equipamentos, pedido cobrindo os dois, histórico com quando/de quem; cabeçalhos novos gravados.")
        _recusa("pelo menos um equipamento", v.criar_venda, fx.CPF_MARIA, ["  "], _pedido("Banco Exemplo"))
        _recusa("Status de proposta", v.criar_venda, fx.CPF_MARIA, ["X"], _pedido("Banco Exemplo", STATUS="Efetivado"))
        _recusa("Informe o banco", v.criar_venda, fx.CPF_MARIA, ["X"], _pedido(""))
        try:
            v.criar_venda("000.000.000-00", ["X"], _pedido("Banco Exemplo"))
            raise AssertionError("cliente não cadastrado deveria ser recusado pela validação de sempre")
        except propostas_mod.ErroProposta:
            pass
        assert len(bd.ler_vendas(caminho)) == 1, "recusa não grava nada"
        print("OK: sem equipamento, sem banco, status de venda numa proposta e cliente não cadastrado são recusados, sem gravar nada.")

        linha("3) Mandar a outro banco (inclusive um pedido só de parte da venda)")
        id_p2 = v.mandar_a_outro_banco(id_venda, _pedido("Outro Banco Exemplo"))
        id_p3 = v.mandar_a_outro_banco(id_venda, _pedido("Banco Terceiro", 20000, EQUIPAMENTO="laser exemplo"))
        da_venda = v.propostas_da_venda(id_venda).set_index("ID_PROPOSTA")
        assert list(da_venda.index) == [id_p1, id_p2, id_p3]
        assert da_venda.loc[id_p3, "EQUIPAMENTO"] == "laser exemplo" and da_venda.loc[id_p2, "EQUIPAMENTO"] == venda["EQUIPAMENTOS"]
        print("OK: 3 propostas na mesma venda; uma delas só com o laser.")

        linha("4) Status da proposta: só os 4 do banco; 'todas negadas' sugere Perdida")
        _recusa("Status de proposta", v.mudar_status_proposta, id_p1, "Nota Fiscal Anexada")
        assert v.mudar_status_proposta(id_p1, propostas_mod.STATUS_NEGADO) is False
        assert v.mudar_status_proposta(id_p2, propostas_mod.STATUS_NEGADO) is False, "ainda falta uma"
        assert v.mudar_status_proposta(id_p3, propostas_mod.STATUS_NEGADO) is True, "a última negada sugere Perdida"
        assert bd.ler_vendas(caminho).iloc[0]["STATUS"] == v.VENDA_AGUARDANDO, "o app nunca marca Perdida sozinho"
        assert v.mudar_status_proposta(id_p3, propostas_mod.STATUS_NEGADO) is False, "sem mudança, nada a sugerir"
        trocas = bd.ler_historico(caminho)
        assert list(trocas[trocas["ID"] == id_p1]["PARA"]) == [propostas_mod.STATUS_EM_ANALISE, propostas_mod.STATUS_NEGADO]
        assert trocas[trocas["ID"] == id_p1]["DE"].iloc[-1] == propostas_mod.STATUS_EM_ANALISE
        print("OK: só Em Análise/Pré-aprovado/Aprovado/Negado; a última negada sugere Perdida (sem marcar sozinho); histórico com de→para.")

        linha("5) Bancos fechados (vários na venda), Efetivado por banco, venda calculada, desfechos e reabrir")
        _recusa("que aprovou", v.fechar_proposta, id_p1)  # negada
        v.mudar_status_proposta(id_p2, propostas_mod.STATUS_APROVADO)  # cobre a venda toda
        v.mudar_status_proposta(id_p3, propostas_mod.STATUS_APROVADO)  # só o laser
        _recusa("marcado como fechado pode ser efetivado", v.efetivar_proposta, id_p2)
        v.fechar_proposta(id_p2)
        venda_ = lambda: bd.ler_vendas(caminho).set_index("ID_VENDA").loc[id_venda]  # noqa: E731
        assert venda_()["STATUS"] == v.VENDA_EM_ANDAMENTO and v.fechadas_da_venda(venda_()["BANCO_ESCOLHIDO"]) == [id_p2]
        _recusa("já está coberto", v.fechar_proposta, id_p3)  # o laser já está no pedido do id_p2
        _recusa("Desmarque o anel", v.mudar_status_proposta, id_p2, propostas_mod.STATUS_NEGADO)
        v.desfazer_fechamento(id_p2)
        assert venda_()["STATUS"] == v.VENDA_AGUARDANDO and venda_()["BANCO_ESCOLHIDO"] == ""
        id_p5 = v.mandar_a_outro_banco(id_venda, _pedido("Banco Quinto", EQUIPAMENTO="Cadeira Exemplo"))
        v.mudar_status_proposta(id_p5, propostas_mod.STATUS_APROVADO)
        v.fechar_proposta(id_p5)  # a cadeira por um banco...
        v.fechar_proposta(id_p3)  # ...e o laser por outro
        assert v.fechadas_da_venda(venda_()["BANCO_ESCOLHIDO"]) == [id_p5, id_p3]
        assert v.efetivar_proposta(id_p5) == v.VENDA_EM_ANDAMENTO, "falta o outro banco"
        _recusa("já está Efetivado", v.desfazer_fechamento, id_p5)
        assert v.efetivar_proposta(id_p3) == v.VENDA_EFETIVADA, "os dois fechados efetivados: a venda está Efetivada"
        hist = bd.ler_historico(caminho)
        assert list(hist[hist["ID"] == id_venda]["PARA"]) == [
            v.VENDA_AGUARDANDO, v.VENDA_EM_ANDAMENTO, v.VENDA_AGUARDANDO, v.VENDA_EM_ANDAMENTO, v.VENDA_EFETIVADA]
        assert set(hist[hist["TIPO"] == v.TIPO_FECHAMENTO]["ID"]) == {id_p2, id_p5, id_p3}
        v.voltar_proposta(id_p3)
        assert venda_()["STATUS"] == v.VENDA_EM_ANDAMENTO, "voltar um banco reabre o andamento"
        assert v.efetivar_proposta(id_p3) == v.VENDA_EFETIVADA
        id_p6 = v.mandar_a_outro_banco(id_venda, _pedido("Banco Sexto"))  # Efetivada ainda aceita outro pedido
        assert venda_()["STATUS"] == v.VENDA_EFETIVADA, "pedido novo (não fechado) não muda a venda"
        v.excluir_proposta(id_p6)
        print("OK: só banco aprovado fecha; vários fechados sem repetir equipamento; cada um efetiva sozinho; a venda é "
              "calculada (Aguardando → Em andamento → Efetivada) e tudo vai pro histórico.")

        outra_venda, outra_p = v.criar_venda(fx.CPF_JOAO, ["Mesa Exemplo"], _pedido("Banco Exemplo"))
        _recusa("vem dos bancos", v.mudar_status_venda, outra_venda, v.VENDA_EFETIVADA)
        _recusa("Motivo inválido", v.mudar_status_venda, outra_venda, v.VENDA_PERDIDA, motivo="Porque sim")
        v.mudar_status_venda(outra_venda, v.VENDA_NAO_EFETIVADA, motivo="Comprou à vista")
        assert bd.ler_vendas(caminho).set_index("ID_VENDA").loc[outra_venda, "MOTIVO"] == "Comprou à vista"
        _recusa("reabra a venda", v.mandar_a_outro_banco, outra_venda, _pedido("Banco Exemplo"))
        v.mudar_status_proposta(outra_p, propostas_mod.STATUS_APROVADO)
        _recusa("reabra a venda", v.fechar_proposta, outra_p)
        assert v.reabrir_venda(outra_venda) == v.VENDA_AGUARDANDO
        assert bd.ler_vendas(caminho).set_index("ID_VENDA").loc[outra_venda, "MOTIVO"] == ""
        _recusa("não há o que reabrir", v.reabrir_venda, outra_venda)
        v.mudar_status_venda(outra_venda, v.VENDA_PERDIDA)  # motivo é opcional
        assert v.voltar_venda(outra_venda) == v.VENDA_AGUARDANDO, "o ↩ da venda reabre"
        # uma venda gravada com as etapas de antes do modelo B ("Nota fiscal"...) aparece com o status dos bancos
        vendas_ = bd.ler_vendas(caminho)
        vendas_.loc[vendas_["ID_VENDA"] == outra_venda, "STATUS"] = "Nota fiscal"
        bd.escrever_vendas_e_propostas(caminho, vendas=vendas_)
        assert next(r for r in v.listar_vendas() if r.id == outra_venda).status == v.VENDA_AGUARDANDO
        print("OK: desfechos manuais com motivo opcional da lista; numa venda desistida/perdida nada fecha nem entra; "
              "reabrir volta ao status calculado.")

        linha("6) Alertas de 7 dias")
        dez_dias = pd.Timestamp(date.today() - timedelta(days=10))
        id_v3, id_p4 = v.criar_venda(fx.CPF_ANA, ["Autoclave Exemplo"], _pedido("Banco Exemplo", DATA=dez_dias))
        hist = bd.ler_historico(caminho)
        hist.loc[hist["ID"].isin([id_v3, id_p4]), "QUANDO"] = dez_dias  # como se tivesse sido criada há 10 dias
        bd.escrever_tudo(caminho, bd.ler_clientes(caminho), bd.ler_equipamentos(caminho), bd.ler_vendedores(caminho),
                         bd.ler_propostas(caminho), historico=hist)
        alertas = {(a.tipo, a.id): a for a in v.alertas()}
        assert (v.TIPO_PROPOSTA, id_p4) in alertas and alertas[(v.TIPO_PROPOSTA, id_p4)].dias == 10
        assert (v.TIPO_VENDA, id_v3) in alertas and "parada" in alertas[(v.TIPO_VENDA, id_v3)].texto
        assert (v.TIPO_VENDA, id_venda) not in alertas, "venda efetivada não alerta"
        assert all(a.id for a in alertas.values()), "propostas antigas sem ID não entram"
        v.mudar_status_proposta(id_p4, propostas_mod.STATUS_PRE_APROVADO)
        assert (v.TIPO_PROPOSTA, id_p4) not in {(a.tipo, a.id) for a in v.alertas()}, "mudou hoje: zera a contagem"
        assert not {(a.tipo, a.id) for a in v.alertas(date.today() - timedelta(days=5))} & {(v.TIPO_VENDA, id_v3)}
        print("OK: proposta sem resposta e venda parada há 10 dias alertam; mudar hoje zera; finalizada não alerta.")

        linha("7) Editar uma proposta pelo caminho antigo mantém o vínculo com a venda")
        indice = int(bd.ler_propostas(caminho).index[bd.ler_propostas(caminho)["ID_PROPOSTA"] == id_p1][0])
        propostas_mod.atualizar_proposta(indice, {"OBSERVAÇÕES": "editada pela tela de sempre"})
        editada = bd.ler_propostas(caminho).loc[indice]
        assert editada["ID_PROPOSTA"] == id_p1 and editada["ID_VENDA"] == id_venda
        print("OK: atualizar_proposta preserva ID_PROPOSTA e ID_VENDA.")

        linha("8) Sincronização: as abas novas vão para a nuvem; nuvem sem elas não apaga as daqui")
        enviados = []
        original = sheets_sync.sincronizar_em_background
        sheets_sync.sincronizar_em_background = lambda aba, df: enviados.append(aba)
        try:
            v.mandar_a_outro_banco(id_v3, _pedido("Outro Banco Exemplo"))
            assert enviados == [bd.ABA_PROPOSTAS, bd.ABA_HISTORICO], enviados
            enviados.clear()
            v.mudar_status_venda(id_v3, v.VENDA_PERDIDA)
            assert enviados == [bd.ABA_VENDAS, bd.ABA_HISTORICO], enviados
            enviados.clear()
            sincronizacao.sincronizar_tudo_agora(caminho)
            assert {bd.ABA_VENDAS, bd.ABA_HISTORICO} <= set(enviados) and len(enviados) == 7, enviados
        finally:
            sheets_sync.sincronizar_em_background = original
        print("OK: cada gravação envia só as abas que mudaram; 'Sincronizar agora' envia as 7.")

        leitores = {n: getattr(nuvem, n) for n in ("ler_clientes", "ler_equipamentos", "ler_vendedores", "ler_propostas",
                                                     "ler_bancos", "ler_vendas", "ler_historico")}
        try:
            for nome, df in (("ler_clientes", bd.ler_clientes(caminho)), ("ler_equipamentos", bd.ler_equipamentos(caminho)),
                             ("ler_vendedores", bd.ler_vendedores(caminho)), ("ler_propostas", bd.ler_propostas(caminho))):
                setattr(nuvem, nome, lambda df=df: df.copy())

            def _sem_aba(nome):
                def _ler():
                    raise gspread.WorksheetNotFound(nome)
                return _ler

            for nome in ("ler_bancos", "ler_vendas", "ler_historico"):
                setattr(nuvem, nome, _sem_aba(nome))
            vendas_antes, hist_antes = bd.ler_vendas(caminho), bd.ler_historico(caminho)
            dados = sincronizacao._ler_dados_da_nuvem()
            assert set(dados) == {bd.ABA_CLIENTES, bd.ABA_EQUIPAMENTOS, bd.ABA_VENDEDORES, bd.ABA_PROPOSTAS}
            sincronizacao._recusar_se_a_nuvem_parece_vazia_por_engano(dados, caminho)
            sincronizacao._gravar_dados_no_arquivo(caminho, dados)
            assert bd.ler_vendas(caminho).equals(vendas_antes) and len(bd.ler_historico(caminho)) == len(hist_antes)
            assert (bd.ler_propostas(caminho)["ID_VENDA"] != "").sum() >= 6, "os vínculos das propostas vêm junto"
        finally:
            for nome, funcao in leitores.items():
                setattr(nuvem, nome, funcao)
        print("OK: nuvem de antes das vendas: baixa o resto e as vendas/histórico daqui ficam iguais.")

        linha("8b) Venda sem proposta, editar a venda, editar e excluir uma proposta")
        _recusa("cliente cadastrado", v.criar_venda, "000.000.000-00", ["Mesa"])
        id_s, sem = v.criar_venda(fx.CPF_ANA, ["Mesa Exemplo"], observacoes="só a venda")
        venda_s = bd.ler_vendas(caminho).set_index("ID_VENDA").loc[id_s]
        assert sem is None and v.propostas_da_venda(id_s).empty and venda_s["STATUS"] == v.VENDA_AGUARDANDO
        assert venda_s["DATA_CRIACAO"] == pd.Timestamp(date.today()) and venda_s["OBSERVAÇÕES"] == "só a venda"
        v.editar_venda(id_s, ["Mesa Exemplo", " cadeira x "], "  outra obs ")
        venda_s = bd.ler_vendas(caminho).set_index("ID_VENDA").loc[id_s]
        assert venda_s["EQUIPAMENTOS"] == "Mesa Exemplo + cadeira x" and venda_s["OBSERVAÇÕES"] == "outra obs"
        _recusa("pelo menos um equipamento", v.editar_venda, id_s, ["  "], "")
        ps = v.mandar_a_outro_banco(id_s, _pedido("Banco Exemplo"))
        hist_antes = len(bd.ler_historico(caminho))
        assert v.editar_proposta(ps, {"VALOR (R$)": 61000, "MESES": 24}) is False
        linha_ps = bd.ler_propostas(caminho).loc[v.indice_da_proposta(ps)]
        assert linha_ps["VALOR (R$)"] == 61000 and linha_ps["MESES"] == 24
        assert len(bd.ler_historico(caminho)) == hist_antes, "sem troca de status, nada no histórico"
        modificado = os.path.getmtime(caminho)
        assert v.editar_proposta(ps, {"VALOR (R$)": 61000}) is False and os.path.getmtime(caminho) == modificado, \
            "nada mudou: não grava"
        _recusa("Status de proposta inválido", v.editar_proposta, ps, {"STATUS": propostas_mod.STATUS_NF_ANEXADA})
        _recusa("marcado como fechado pode ser efetivado", v.editar_proposta, ps, {"STATUS": propostas_mod.STATUS_EFETIVADO})
        aberta = bd.ler_propostas(caminho).loc[v.indice_da_proposta(ps)].to_dict()  # como o formulario abriu
        v.editar_proposta(ps, {"OBSERVAÇÕES": "outra tela mudou"})
        try:
            v.editar_proposta(ps, {"MESES": 12}, esperado=aberta)
            raise AssertionError("a proposta mudou desde que foi aberta: deveria recusar")
        except propostas_mod.ErroProposta as exc:
            assert "alterada ou excluída" in str(exc)
        try:
            v.editar_proposta(ps, {"VALOR (R$)": -5})
            raise AssertionError("valor negativo deveria ser recusado")
        except propostas_mod.ErroProposta:
            pass
        assert v.editar_proposta(ps, {"STATUS": propostas_mod.STATUS_NEGADO}) is True, "a única negou: sugere Perdida"
        hist = bd.ler_historico(caminho)
        assert ((hist["ID"] == ps) & (hist["PARA"] == propostas_mod.STATUS_NEGADO)).any()
        v.editar_proposta(ps, {"STATUS": propostas_mod.STATUS_APROVADO})
        v.fechar_proposta(ps)
        _recusa("Desmarque o anel", v.editar_proposta, ps, {"STATUS": propostas_mod.STATUS_NEGADO})
        try:
            v.excluir_proposta(ps)
            raise AssertionError("um banco fechado não pode ser excluído")
        except propostas_mod.ErroProposta as exc:
            assert "fechado" in str(exc)
        v.desfazer_fechamento(ps)
        v.excluir_proposta(ps)
        assert v.propostas_da_venda(id_s).empty
        print("OK: venda sem proposta (cliente cadastrado); editar a venda; editar a proposta com as regras e o histórico, "
              "sem gravar quando nada muda; excluir (menos um banco fechado).")

        linha("9) Só o Administrador grava")
        sessao_mod.encerrar()
        try:
            v.criar_venda(fx.CPF_MARIA, ["X"], _pedido("Banco Exemplo"))
            raise AssertionError("sem sessão não pode gravar")
        except sessao_mod.PermissaoNegada:
            pass
        print("OK: sem sessão de Administrador, criar venda é recusado.")
        linha("TUDO OK")
    finally:
        nuvem._obter_cliente, sheets_sync._obter_cliente = rede_original
        fx.restaurar_modulos()
        sessao_mod.encerrar()


if __name__ == "__main__":
    main()
