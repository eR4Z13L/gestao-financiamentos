"""Testa a separacao de endereco (core/endereco.py) e as migracoes da aba
CLIENTES (core/migracao_endereco.py) a partir dos DOIS formatos antigos
(legado: "ENDEREÇO" unico; intermediario: campos separados sem COMPLEMENTO/UF)
- so com dados FICTICIOS, em planilhas montadas aqui mesmo em pasta
temporaria. Nao mexe na planilha real.

Rodar com: venv/Scripts/python.exe scripts/smoke_test_migracao_endereco.py
"""

from __future__ import annotations

import hashlib
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
config.SINCRONIZACAO_GOOGLE_ATIVADA = False  # nunca manda dado de teste pra planilha real na nuvem
from core import clientes as clientes_mod
from core import data_store as bd
from core import migracao_endereco as mig
from core import propostas as propostas_mod
from core import sessao as sessao_mod
from core.endereco import separar_cidade_uf, separar_endereco


def linha(titulo: str) -> None:
    print(f"\n{'=' * 60}\n{titulo}\n{'=' * 60}")


def _hash(caminho: Path) -> str:
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


# ---------------------------------------------------------------------------
# 1) parser
# ---------------------------------------------------------------------------

CASOS_SEGUROS = [
    # (texto, cep, logradouro, numero, complemento, bairro, cidade, uf)
    ("Rua das Flores, 123 - Centro/Curitiba - PR 80000-000", "80000-000", "Rua das Flores", "123", "", "Centro", "Curitiba", "PR"),
    ("AV BRASIL, SN - JARDIM /SANTOS SP CEP 11000000", "11000-000", "AV BRASIL", "SN", "", "JARDIM", "SANTOS", "SP"),
    ("Rua A, 45A - Vila Nova - Recife - PE", "", "Rua A", "45A", "", "Vila Nova", "Recife", "PE"),
    ("Rua Sem Cep, 10. Bairro X. Cidade Y/MG", "", "Rua Sem Cep", "10", "", "Bairro X", "Cidade Y", "MG"),
    ("Rua B, 5 - Sé/São Paulo - SP 01001-000", "01001-000", "Rua B", "5", "", "Sé", "São Paulo", "SP"),  # CEP com zero na frente
    ("Rua A,\n 12 -  Centro / Natal - RN\n", "", "Rua A", "12", "", "Centro", "Natal", "RN"),  # espacos e quebra de linha
    ("R. DOS TESTES, 2697 JD MODELO - CIDADE FICTICIA - SP -  CEP 15503-443", "15503-443", "R. DOS TESTES", "2697", "", "JD MODELO", "CIDADE FICTICIA", "SP"),
    ("Rua C, S/N - Centro/Belem - PA cep: 66000000", "66000-000", "Rua C", "S/N", "", "Centro", "Belem", "PA"),
    # com complemento no formato reconhecido ("palavra + identificador" logo depois do numero)
    ("Rua X, 12 apto 3 - Centro/Natal - RN", "", "Rua X", "12", "apto 3", "Centro", "Natal", "RN"),
    ("Rua X, 12 - Bloco B - Centro/Natal - RN", "", "Rua X", "12", "Bloco B", "Centro", "Natal", "RN"),
    ("Rua X, 12 BLOCO B APTO 12 - Centro/Natal - RN", "", "Rua X", "12", "BLOCO B APTO 12", "Centro", "Natal", "RN"),
    ("Rua X, 12, sala 5, Centro/Natal - RN", "", "Rua X", "12", "sala 5", "Centro", "Natal", "RN"),
    ("Rua Z, 374 APTO 506 SANTA CRUZ/GUARAPUAVA - PR 85015080", "85015-080", "Rua Z", "374", "APTO 506", "SANTA CRUZ", "GUARAPUAVA", "PR"),
    ("Rua Q, 8 casa 2 - Centro/Natal - RN", "", "Rua Q", "8", "casa 2", "Centro", "Natal", "RN"),
]

CASOS_RECUSADOS = [
    # (texto, trecho esperado no motivo)
    ("", "vazio"),
    ("   \n  ", "vazio"),
    ("mora logo ali", "padrão"),
    ("Rua X 123", "padrão"),  # sem virgula antes do numero
    ("Rua X, Nº 12 - Centro/Natal - RN", "padrão"),  # "Nº" nao e o numero puro
    ("Rua X, 12 - Casa Amarela/Recife - PE", "complemento"),  # "casa" sem identificador: conservador, revisao
    ("Rua X, 12 C CS APTO C - Centro/Natal - RN", "complemento"),  # formato de complemento desconhecido
    ("Rua X, 12 apto - Centro/Natal - RN", "complemento"),  # palavra de complemento sem identificador
    ("Rua X, 12 - Centro - Natal", "UF"),
    ("Rua X, 12 - Centro - Natal - Bahia", "UF"),  # estado por extenso
    ("Rua X, 12 - Centro/Natal - XX", "UF"),  # sigla inexistente
    ("Rua X, 12 Centro Natal/RN", "1 bloco"),  # sem separador entre bairro e cidade
    ("Rua X, 12 apto 3 Centro Natal/RN", "1 bloco"),  # complemento ok, mas bairro/cidade ambiguos
    ("Rua X, 12 - Centro - Zona Sul/Natal - RN", "3 bloco"),  # bairro/cidade ambiguos
    ("Rua X, 12 - Centro/Natal - RN 59000-000 59000-001", "mais de um CEP"),
]


def testar_parser() -> None:
    linha("1) core/endereco.py - separa so o que segue padrao claro")
    for texto, cep, logr, num, compl, bairro, cidade, uf in CASOS_SEGUROS:
        resultado, motivo = separar_endereco(texto)
        assert resultado is not None, f"deveria separar {texto!r}, mas recusou: {motivo}"
        obtido = (resultado.cep, resultado.logradouro, resultado.numero, resultado.complemento,
                  resultado.bairro, resultado.cidade, resultado.uf)
        assert obtido == (cep, logr, num, compl, bairro, cidade, uf), f"{texto!r} -> {obtido}"
    print(f"OK: {len(CASOS_SEGUROS)} endereços de padrão claro separados corretamente (com e sem complemento).")

    for texto, trecho in CASOS_RECUSADOS:
        resultado, motivo = separar_endereco(texto)
        assert resultado is None, f"NAO deveria separar {texto!r}, mas separou: {resultado}"
        assert trecho in motivo, f"{texto!r}: motivo {motivo!r} deveria citar {trecho!r}"
    print(f"OK: {len(CASOS_RECUSADOS)} endereços ambíguos/fora do padrão recusados, cada um com o motivo certo.")

    assert separar_cidade_uf("Fortaleza/CE") == ("Fortaleza", "CE")
    assert separar_cidade_uf("São Paulo / SP") == ("São Paulo", "SP")
    for texto in ("Cidade Solta", "X/ZZ", "Fortaleza/ce", "", "Fortaleza - CE"):
        assert separar_cidade_uf(texto) == (texto, ""), f"{texto!r} nao pode ganhar UF adivinhada"
    print("OK: 'Cidade/UF' é dividida só com UF válida; qualquer outra coisa fica intacta (UF em branco).")


# ---------------------------------------------------------------------------
# 2) planilhas "antigas" ficticias
# ---------------------------------------------------------------------------

ENDERECO_OK = "Rua das Flores, 123 - Centro/Curitiba - PR 80000-000"
ENDERECO_DUVIDOSO = "Rua Sem Padrao 77 perto do mercado"
ENDERECO_COM_APTO = "Rua Z, 9 apto 4 - Centro/Natal - RN"
ENDERECO_COMPLEMENTO_ESTRANHO = "Rua Z, 9 C CS APTO C - Centro/Natal - RN"


def _criar_planilha(caminho: Path, colunas: list[str], clientes: list[dict]) -> None:
    """Monta uma planilha completa (4 abas) com a aba CLIENTES no formato `colunas`."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = bd.ABA_CLIENTES
    ws.append(colunas)
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFFFF")
        celula.fill = PatternFill("solid", fgColor="FF2F5597")
        celula.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 26.25
    ws.freeze_panes = "A2"
    for j in range(1, len(colunas) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 20
    for i, cliente in enumerate(clientes, start=2):
        ws.append([cliente.get(c) for c in colunas])
        for coluna in ("DATA CADASTRO", "NASCIMENTO"):
            ws.cell(row=i, column=colunas.index(coluna) + 1).number_format = "dd/mm/yyyy"

    ws_eq = wb.create_sheet(bd.ABA_EQUIPAMENTOS)
    ws_eq.append(bd.EQUIPAMENTOS_COLUNAS)
    ws_eq.append(["Fornecedor Teste", "Equip Teste", 36, 1200.5, 30000, "obs"])

    ws_prop = wb.create_sheet(bd.ABA_PROPOSTAS)
    ws_prop.append(bd.PROPOSTAS_COLUNAS)
    bd._escrever_linhas_propostas(
        ws_prop,
        [
            {"DATA": datetime(2026, 2, 1), "CPF": "52998224725", "VALOR (R$)": 25000, "MESES": 36,
             "EQUIPAMENTO": "Equip Teste", "BANCO": "Banco Teste", "STATUS": "Em Análise", "OBSERVAÇÕES": "obs prop"},
        ],
    )

    ws_v = wb.create_sheet(bd.ABA_VENDEDORES)
    ws_v.append(bd.VENDEDORES_COLUNAS)
    ws_v.append(["Ana", "", ""])
    ws_v.append(["Bia", "", ""])
    wb.save(caminho)


def _base(cpf: str, vendedor: str, nome: str, **extra) -> dict:
    return {"DATA CADASTRO": datetime(2026, 1, 5), "CPF/CNPJ": cpf, "VENDEDOR": vendedor, "TIPO": "Cliente",
            "CLIENTE": nome, "NASCIMENTO": datetime(1985, 3, 15), "CELULAR": "(85) 99999-0001",
            "VINCULADO": "", "REDE SOCIAL": "@x", "EMAIL": "x@exemplo.com", **extra}


def _planilha_legado(caminho: Path) -> None:
    _criar_planilha(caminho, mig.COLUNAS_LEGADO, [
        _base("52998224725", "Ana", "CLIENTE UM", **{"ENDEREÇO": ENDERECO_OK}),
        _base("11144477735", "Ana", "CLIENTE DOIS", **{"ENDEREÇO": ENDERECO_DUVIDOSO}),
        _base("39053344705", "Bia", "CLIENTE TRES", TIPO="Avalista", VINCULADO="CLIENTE UM", NASCIMENTO=None),
        _base("12345678909", "Bia", "CLIENTE QUATRO", **{"ENDEREÇO": ENDERECO_COM_APTO}),
        _base("00000000191", "Ana", "CLIENTE CINCO", **{"ENDEREÇO": ENDERECO_OK}),
        _base("98765432100", "Bia", "CLIENTE SEIS", **{"ENDEREÇO": ENDERECO_COMPLEMENTO_ESTRANHO}),
    ])


def _planilha_intermediaria(caminho: Path) -> None:
    _criar_planilha(caminho, mig.COLUNAS_INTERMEDIARIO, [
        # 2: campos ja separados, "Cidade/UF" grudada -> divide; pai/mae/profissao ja preenchidos -> preserva
        _base("52998224725", "Ana", "CLIENTE ALFA", CEP="80000-000", LOGRADOURO="Rua das Flores", **{"NÚMERO": "123"},
              BAIRRO="Centro", CIDADE="Curitiba/PR", **{"NOME DO PAI": "Pai Alfa", "PROFISSÃO": "Engenheiro"}),
        # 3: so o texto em REVISAR, com complemento reconhecivel -> o parser atual consegue
        _base("11144477735", "Ana", "CLIENTE BETA", **{"ENDEREÇO (REVISAR)": "Rua Y, 5 apto 2 - Centro/Natal - RN 59000-000"}),
        # 4: so o texto em REVISAR, sem padrao -> continua em revisao
        _base("39053344705", "Bia", "CLIENTE GAMA", **{"ENDEREÇO (REVISAR)": "mora logo ali"}),
        # 5: revisao PELA METADE (ja preencheu logradouro e cidade sem UF) -> nao pode ser tocado
        _base("12345678909", "Bia", "CLIENTE DELTA", LOGRADOURO="Rua Manual", CIDADE="Cidade Solta",
              **{"ENDEREÇO (REVISAR)": "texto antigo"}),
        # 6: sem endereco nenhum
        _base("00000000191", "Ana", "CLIENTE EPSILON"),
        # 7: outra cidade com UF valida
        _base("98765432100", "Bia", "CLIENTE ZETA", LOGRADOURO="Rua Z", CIDADE="Rio de Janeiro/RJ"),
    ])


def _backups(pasta: Path, rotulo: str) -> list[Path]:
    return sorted(pasta.glob(f"*.backup-{rotulo}-*.xlsx"))


def _colunas_de_endereco(cliente) -> tuple:
    return tuple(cliente[c] for c in ("CEP", "LOGRADOURO", "NÚMERO", "COMPLEMENTO", "BAIRRO", "CIDADE", "UF"))


def testar_migracao_legado() -> None:
    linha("2) Migração a partir do formato LEGADO (ENDEREÇO único)")
    pasta = Path(tempfile.mkdtemp(prefix="_smoke_migracao_"))
    try:
        arquivo = pasta / "controle.xlsx"
        _planilha_legado(arquivo)
        hash_original = _hash(arquivo)

        # app recusa ler/gravar o formato antigo em vez de misturar colunas
        try:
            bd.ler_clientes(arquivo)
            raise AssertionError("ler_clientes deveria recusar a planilha no formato antigo")
        except bd.ErroPlanilhaDesatualizada as exc:
            assert "migrar_endereco_clientes.py" in str(exc), str(exc)
        try:
            bd.escrever_clientes(arquivo, bd.ler_aba(arquivo, bd.ABA_CLIENTES, mig.COLUNAS_LEGADO))
            raise AssertionError("escrever_clientes deveria recusar a planilha no formato antigo")
        except bd.ErroPlanilhaDesatualizada:
            pass
        assert _hash(arquivo) == hash_original, "recusar nao pode ter alterado o arquivo"
        print("OK: formato antigo é recusado na leitura e na escrita (mensagem aponta o script), arquivo intacto.")

        # simulacao: classifica sem gravar nada
        _, plano = mig.planejar(arquivo)
        assert _hash(arquivo) == hash_original and not _backups(pasta, "pre-endereco"), "simular nao grava nem cria backup"
        assert plano.formato == mig.FORMATO_LEGADO
        assert (plano.total_clientes, plano.sem_endereco, len(plano.separados), len(plano.revisar)) == (6, 1, 3, 2)
        assert {i.linha_planilha for i in plano.separados} == {2, 5, 6}
        assert {i.linha_planilha for i in plano.revisar} == {3, 7}
        print("OK: simulação classifica (6 clientes: 1 sem endereço, 3 separados, 2 a revisar) sem gravar nada.")

        # arquivo aberto no Excel -> recusa antes de qualquer backup
        trava = pasta / "~$controle.xlsx"
        trava.write_text("")
        try:
            mig.migrar_planilha(arquivo)
            raise AssertionError("deveria recusar com o arquivo aberto no Excel")
        except bd.ErroArquivoBloqueado:
            pass
        assert not _backups(pasta, "pre-endereco") and _hash(arquivo) == hash_original
        trava.unlink()
        print("OK: arquivo aberto no Excel -> recusa sem criar backup nem alterar nada.")

        # conferencia falhando -> o arquivo real nunca e tocado
        original_verificar = mig._verificar_resultado

        def _falha(*_args):
            raise mig.ErroMigracao("falha simulada na conferência")

        mig._verificar_resultado = _falha
        try:
            mig.migrar_planilha(arquivo)
            raise AssertionError("deveria abortar quando a conferência falha")
        except mig.ErroMigracao:
            pass
        finally:
            mig._verificar_resultado = original_verificar
        assert _hash(arquivo) == hash_original, "conferencia falhou: arquivo real nao pode ter mudado"
        assert not list(pasta.glob("*.migrando.xlsx")), "arquivo temporario deveria ter sido apagado"
        for b in _backups(pasta, "pre-endereco"):
            b.unlink()
        print("OK: se a conferência falha, o arquivo real fica intacto e o temporário é apagado.")

        # migracao de verdade
        relatorio = mig.migrar_planilha(arquivo)
        backups = _backups(pasta, "pre-endereco")
        assert len(backups) == 1 and relatorio.caminho_backup == backups[0]
        assert _hash(backups[0]) == hash_original, "backup tem que ser identico ao original"
        print(f"OK: backup criado e idêntico ao original ({backups[0].name}).")

        wb = openpyxl.load_workbook(arquivo)
        ws = wb[bd.ABA_CLIENTES]
        assert [c.value for c in ws[1]] == bd.CLIENTES_COLUNAS
        assert bd.CLIENTES_COLUNAS[:5] == ["DATA CADASTRO", "CPF/CNPJ", "VENDEDOR", "TIPO", "CLIENTE"], "A-E nao podem mudar (VLOOKUP)"
        assert len(bd.CLIENTES_COLUNAS) == 21
        cab = ws["H1"]
        assert cab.font.b and cab.fill.fgColor.rgb == "FF2F5597" and cab.alignment.wrap_text, "cabeçalho novo mantém o estilo"
        assert ws.freeze_panes == "A2" and ws.row_dimensions[1].height == 26.25
        assert all(ws.column_dimensions[get_column_letter(j)].width for j in range(1, 22))
        col = {n: j for j, n in enumerate(bd.CLIENTES_COLUNAS, start=1)}
        assert ws.cell(row=2, column=col["NASCIMENTO"]).number_format == "dd/mm/yyyy"
        assert ws.cell(row=2, column=col["CEP"]).number_format == "@"
        assert ws.cell(row=2, column=col["NÚMERO"]).number_format == "@"
        wb.close()
        print("OK: cabeçalho novo (21 colunas, A-E intactas), estilo/largura/formatos aplicados.")

        df = bd.ler_clientes(arquivo)
        assert len(df) == 6
        um, dois, tres, quatro, cinco, seis = (df.iloc[i] for i in range(6))
        assert _colunas_de_endereco(um) == ("80000-000", "Rua das Flores", "123", "", "Centro", "Curitiba", "PR")
        assert _colunas_de_endereco(cinco) == _colunas_de_endereco(um)
        assert _colunas_de_endereco(quatro) == ("", "Rua Z", "9", "apto 4", "Centro", "Natal", "RN")
        for duvidoso, original in ((dois, ENDERECO_DUVIDOSO), (seis, ENDERECO_COMPLEMENTO_ESTRANHO)):
            assert duvidoso["ENDEREÇO (REVISAR)"] == original, "texto original inteiro em REVISAR"
            assert not any(_colunas_de_endereco(duvidoso)), "nada adivinhado"
        assert not any(_colunas_de_endereco(tres)) and tres["ENDEREÇO (REVISAR)"] == ""
        for _, cliente in df.iterrows():
            assert cliente["NOME DO PAI"] == cliente["NOME DA MÃE"] == cliente["PROFISSÃO"] == "", "campos novos em branco"
        assert (um["CLIENTE"], um["EMAIL"], um["CELULAR"]) == ("CLIENTE UM", "x@exemplo.com", "(85) 99999-0001")
        assert tres["VINCULADO"] == "CLIENTE UM" and tres["TIPO"] == "Avalista"
        assert str(um["NASCIMENTO"])[:10] == "1985-03-15" and str(quatro["DATA CADASTRO"])[:10] == "2026-01-05"
        print("OK: 3 separados (um com complemento), 2 com texto original em 'ENDEREÇO (REVISAR)' (nada adivinhado), "
              "1 sem endereço; pai/mãe/profissão em branco; demais colunas iguais.")

        # PROPOSTAS: formulas iguais e cruzamento por CPF continua funcionando
        propostas = bd.ler_propostas(arquivo)
        assert propostas.iloc[0]["CLIENTE"] == "CLIENTE UM" and propostas.iloc[0]["VENDEDOR"] == "Ana"
        wb = openpyxl.load_workbook(arquivo)
        assert "CLIENTES!$B:$C" in wb[bd.ABA_PROPOSTAS]["B2"].value and "CLIENTES!$B:$E" in wb[bd.ABA_PROPOSTAS]["D2"].value
        wb.close()
        print("OK: aba PROPOSTAS intacta (fórmulas VLOOKUP e cruzamento VENDEDOR/CLIENTE continuam certos).")

        # rodar de novo nao faz nada
        try:
            mig.migrar_planilha(arquivo)
            raise AssertionError("segunda migracao deveria ser recusada")
        except mig.ErroMigracao as exc:
            assert "já está no formato atual" in str(exc)
        assert len(_backups(pasta, "pre-endereco")) == 1 and not _backups(pasta, "pre-complemento-uf")
        print("OK: rodar a migração de novo é recusado (sem novo backup).")

        # o app grava/edita normalmente no formato novo (21 colunas)
        sessao_mod.iniciar(sessao_mod.Sessao(papel=sessao_mod.PAPEL_ADMIN, nome_usuario="Administrador"))
        clientes_mod.CAMINHO_XLSX = arquivo
        propostas_mod.CAMINHO_XLSX = arquivo
        try:
            clientes_mod.atualizar_cliente(
                "39053344705",
                {"CPF/CNPJ": "39053344705", "CLIENTE": "CLIENTE TRES", "TIPO": "Avalista",
                 "CEP": "02378-255", "LOGRADOURO": "Rua Nova", "NÚMERO": "8", "COMPLEMENTO": "sala 5", "BAIRRO": "Centro",
                 "CIDADE": "Natal", "UF": "rn",
                 "NOME DO PAI": "Pai Teste", "NOME DA MÃE": "Mãe Teste", "PROFISSÃO": "Engenheiro"},
            )
            try:
                clientes_mod.atualizar_cliente(
                    "39053344705", {"CPF/CNPJ": "39053344705", "CLIENTE": "CLIENTE TRES", "TIPO": "Avalista", "UF": "Ceara"})
                raise AssertionError("UF por extenso deveria ser recusada")
            except clientes_mod.ErroCliente as exc:
                assert "UF" in str(exc)
        finally:
            clientes_mod.CAMINHO_XLSX = config.CAMINHO_XLSX
            propostas_mod.CAMINHO_XLSX = config.CAMINHO_XLSX
            sessao_mod.encerrar()
        df2 = bd.ler_clientes(arquivo)
        tres2 = df2[df2["CPF/CNPJ"] == "39053344705"].iloc[0]
        assert _colunas_de_endereco(tres2) == ("02378-255", "Rua Nova", "8", "sala 5", "Centro", "Natal", "RN"), "UF 'rn' vira 'RN'"
        assert (tres2["NOME DO PAI"], tres2["NOME DA MÃE"], tres2["PROFISSÃO"]) == ("Pai Teste", "Mãe Teste", "Engenheiro")
        assert df2.iloc[0]["LOGRADOURO"] == "Rua das Flores" and df2.iloc[1]["ENDEREÇO (REVISAR)"] == ENDERECO_DUVIDOSO
        print("OK: depois de migrada, o app edita normalmente (CEP com zero na frente preservado, UF normalizada, "
              "UF inválida recusada; outras linhas intactas).")
    finally:
        shutil.rmtree(pasta, ignore_errors=True)


def testar_migracao_intermediario() -> None:
    linha("3) Migração a partir do formato INTERMEDIÁRIO (19 colunas: sem COMPLEMENTO, UF junto da cidade)")
    pasta = Path(tempfile.mkdtemp(prefix="_smoke_migracao2_"))
    try:
        arquivo = pasta / "controle.xlsx"
        _planilha_intermediaria(arquivo)
        hash_original = _hash(arquivo)

        try:
            bd.ler_clientes(arquivo)
            raise AssertionError("o formato intermediario tambem tem que ser recusado pelo app")
        except bd.ErroPlanilhaDesatualizada:
            pass

        _, plano = mig.planejar(arquivo)
        assert _hash(arquivo) == hash_original
        assert plano.formato == mig.FORMATO_INTERMEDIARIO
        assert plano.total_clientes == 6 and plano.sem_endereco == 1 and plano.ja_separados == 3
        assert plano.cidades_divididas == 2 and [c[0] for c in plano.cidades_sem_uf] == [5]
        assert [i.linha_planilha for i in plano.separados] == [3] and [i.linha_planilha for i in plano.revisar] == [4]
        print("OK: simulação: 2 cidades divididas, 1 endereço em revisão agora separado, 1 continua em revisão, "
              "1 cidade sem UF reconhecível listada.")

        relatorio = mig.migrar_planilha(arquivo)
        backups = _backups(pasta, "pre-complemento-uf")
        assert len(backups) == 1 and _hash(backups[0]) == hash_original and relatorio.caminho_backup == backups[0]
        print(f"OK: backup criado e idêntico ({backups[0].name}).")

        df = bd.ler_clientes(arquivo)
        alfa, beta, gama, delta, epsilon, zeta = (df.iloc[i] for i in range(6))
        assert _colunas_de_endereco(alfa) == ("80000-000", "Rua das Flores", "123", "", "Centro", "Curitiba", "PR")
        assert (alfa["NOME DO PAI"], alfa["PROFISSÃO"]) == ("Pai Alfa", "Engenheiro"), "campos ja preenchidos preservados"
        assert _colunas_de_endereco(beta) == ("59000-000", "Rua Y", "5", "apto 2", "Centro", "Natal", "RN")
        assert beta["ENDEREÇO (REVISAR)"] == "", "separado com sucesso: o texto de revisão é limpo"
        assert gama["ENDEREÇO (REVISAR)"] == "mora logo ali" and not any(_colunas_de_endereco(gama))
        assert _colunas_de_endereco(delta) == ("", "Rua Manual", "", "", "", "Cidade Solta", ""), "revisão pela metade: intacta"
        assert delta["ENDEREÇO (REVISAR)"] == "texto antigo"
        assert not any(_colunas_de_endereco(epsilon)) and epsilon["ENDEREÇO (REVISAR)"] == ""
        assert (zeta["CIDADE"], zeta["UF"]) == ("Rio de Janeiro", "RJ")
        print("OK: Cidade/UF divididas; revisão que agora dá certo é separada e limpa; revisão pela metade e "
              "cidade sem UF ficam exatamente como estavam; dados preenchidos preservados.")

        try:
            mig.migrar_planilha(arquivo)
            raise AssertionError("segunda migracao deveria ser recusada")
        except mig.ErroMigracao as exc:
            assert "já está no formato atual" in str(exc)
        print("OK: rodar de novo é recusado.")
    finally:
        shutil.rmtree(pasta, ignore_errors=True)


if __name__ == "__main__":
    testar_parser()
    testar_migracao_legado()
    testar_migracao_intermediario()
    linha("TUDO OK")
