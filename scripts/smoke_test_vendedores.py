"""Valida o cadastro de vendedores (core/vendedores.py + core/data_store.py):
1. Migracao automatica: um arquivo que ainda nao tem a aba VENDEDORES ganha
   uma na primeira leitura, populada com os nomes distintos ja usados em
   CLIENTES (sem duplicar por maiusculas/espacos).
2. Cadastro novo nao duplica por diferenca de maiusculas/espacos.
3. Nome vazio e rejeitado.

Nunca mexe no arquivo real - tudo numa copia temporaria.

Rodar com: venv/Scripts/python.exe scripts/smoke_test_vendedores.py
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from config import CAMINHO_XLSX

import config
config.SINCRONIZACAO_GOOGLE_ATIVADA = False  # nunca manda dado de teste pra planilha real na nuvem
from core import data_store as bd
from core import sessao as sessao_mod
from core import vendedores as vendedores_mod


def linha(titulo: str) -> None:
    print(f"\n{'=' * 60}\n{titulo}\n{'=' * 60}")


def testar_migracao_automatica(tmp_path: Path) -> None:
    linha("1) Migração automática da aba VENDEDORES")

    wb = openpyxl.load_workbook(tmp_path)
    assert bd.ABA_VENDEDORES not in wb.sheetnames, "arquivo de teste ja tinha a aba - copia nao é mais 'antiga'"
    clientes = bd.ler_clientes(tmp_path)
    esperados = sorted({v.upper() for v in clientes["VENDEDOR"] if v})
    wb.close()

    vendedores = bd.ler_vendedores(tmp_path)
    encontrados = sorted({v.upper() for v in vendedores["NOME"] if v})
    print(f"vendedores distintos em CLIENTES: {len(esperados)} | migrados pra VENDEDORES: {len(encontrados)}")
    assert encontrados == esperados, "migração deveria trazer exatamente os nomes distintos (case-insensitive) de CLIENTES"

    wb2 = openpyxl.load_workbook(tmp_path)
    assert bd.ABA_VENDEDORES in wb2.sheetnames, "migração deveria ter salvo a aba nova no arquivo"
    wb2.close()
    print("OK: aba VENDEDORES criada e persistida na primeira leitura, com os nomes já em uso.")

    n_linhas_apos_migrar = len(bd.ler_vendedores(tmp_path))
    assert n_linhas_apos_migrar == len(vendedores), "reler nao deveria migrar de novo nem duplicar linhas"
    print("OK: reler não migra de novo (idempotente).")


def testar_cadastro_sem_duplicar(tmp_path: Path) -> None:
    linha("2) Cadastro de vendedor novo - sem duplicar por maiúsculas/espaço")

    vendedores_mod.CAMINHO_XLSX = tmp_path
    antes = len(vendedores_mod.listar_vendedores())

    salvo = vendedores_mod.adicionar_vendedor("  Vendedor Smoke Novo  ")
    assert salvo == "Vendedor Smoke Novo", f"deveria remover espaço nas pontas, veio {salvo!r}"
    assert "Vendedor Smoke Novo" in vendedores_mod.listar_vendedores()
    depois_do_primeiro = len(vendedores_mod.listar_vendedores())
    assert depois_do_primeiro == antes + 1
    print(f"OK: cadastrou 'Vendedor Smoke Novo' ({antes} -> {depois_do_primeiro}).")

    salvo_de_novo = vendedores_mod.adicionar_vendedor("VENDEDOR smoke NOVO")
    assert salvo_de_novo == "Vendedor Smoke Novo", "deveria devolver a grafia já cadastrada, não criar outra"
    assert len(vendedores_mod.listar_vendedores()) == depois_do_primeiro, "não deveria ter duplicado por causa/espaço"
    print("OK: cadastrar de novo com outra caixa/espaço não duplica - devolve o nome já existente.")

    try:
        vendedores_mod.adicionar_vendedor("   ")
        raise SystemExit("deveria ter rejeitado nome vazio")
    except vendedores_mod.ErroVendedor as exc:
        print(f"OK: nome vazio rejeitado -> {exc}")


def main() -> None:
    sessao_mod.iniciar(sessao_mod.Sessao(papel=sessao_mod.PAPEL_ADMIN, nome_usuario="Administrador"))
    tmp_path = CAMINHO_XLSX.parent / "_smoke_test_vendedores.xlsx"
    shutil.copy(CAMINHO_XLSX, tmp_path)
    # o arquivo real ja passou pela migracao (a aba VENDEDORES foi criada nele
    # numa sessao anterior) - remove a aba da COPIA antes do teste, pra
    # simular de proposito um arquivo "antigo" (pre-migracao), em vez de
    # depender do arquivo real por acaso ainda nao ter a aba.
    wb = openpyxl.load_workbook(tmp_path)
    if bd.ABA_VENDEDORES in wb.sheetnames:
        del wb[bd.ABA_VENDEDORES]
        wb.save(tmp_path)
    wb.close()
    try:
        testar_migracao_automatica(tmp_path)
        testar_cadastro_sem_duplicar(tmp_path)
        linha("TUDO OK")
    finally:
        vendedores_mod.CAMINHO_XLSX = CAMINHO_XLSX
        tmp_path.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
