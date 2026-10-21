"""Testa as pendencias da venda em qualquer etapa (plano 1.2-B, parte 1: regra em core/):

- varias pendencias por venda (texto, data de criacao, data de resolucao), adicionar e resolver, ordem, ids;
- texto vazio e recusado; resolver outra vez a mesma, uma que nao existe ou numa venda que nao existe: recusa com
  mensagem e nada e gravado; so o ADMIN escreve;
- cada criacao e cada resolucao grava uma linha de TIPO "pendencia" no HISTORICO (ID = a venda);
- pendencia nao e status e nao trava nada (mandar a outro banco, fechar, mudar etapa, desfechos) e NUNCA conta para a
  parada de 7 dias: nem gera alerta/selo, nem zera a parada de uma venda que de fato esta parada;
- ida e volta com a nuvem simulada; planilha de antes da coluna PENDENCIAS abre sem erro e ganha a coluna ao gravar;
  celula ilegivel e recusada com mensagem;
- "NF enviada" (nome novo) e "Nota fiscal enviada" (nome antigo, lido como sinonimo).

Planilha 100% ficticia em pasta temporaria (scripts/fixture_ficticia.py), preferencias isoladas e o Google desligado.

Rodar com: venv/Scripts/python.exe scripts/smoke_test_pendencias.py
"""

from __future__ import annotations

import hashlib
import os
import sys
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl
import pandas as pd

import config

config.SINCRONIZACAO_GOOGLE_ATIVADA = False
import fixture_ficticia as fx
from core import data_store as bd
from core import data_store_sheets as nuvem
from core import migracao_legado
from core import propostas as propostas_mod
from core import resumo_vendas
from core import sessao as sessao_mod
from core import sheets_sync
from core import vendas as v

ADMIN = sessao_mod.Sessao(papel=sessao_mod.PAPEL_ADMIN, nome_usuario="Administrador")
APROVADO, ASSINADO, NOTA, EFETIVADO = (propostas_mod.STATUS_APROVADO, propostas_mod.STATUS_ASSINADO,
                                       propostas_mod.STATUS_NOTA_ENVIADA, propostas_mod.STATUS_EFETIVADO)


def linha(titulo: str) -> None:
    print(f"\n{'=' * 60}\n{titulo}\n{'=' * 60}")


def _recusa(trecho: str, funcao, *args, **kwargs) -> None:
    try:
        funcao(*args, **kwargs)
    except v.ErroVenda as exc:
        assert trecho in str(exc), (trecho, str(exc))
        return
    raise AssertionError(f"deveria recusar ({trecho})")


def _pedido(banco: str, valor: float = 50000, **extra) -> dict:
    return {"BANCO": banco, "VALOR (R$)": valor, "MESES": 36, **extra}


class _Relogio(datetime):
    """O datetime de core/vendas.py com a hora combinada: o historico e as datas das pendencias saem dela."""
    agora_combinado = datetime(2026, 1, 1)

    @classmethod
    def now(cls, tz=None):
        return cls.agora_combinado


def main() -> None:
    def _sem_rede(*_a, **_k):
        raise AssertionError("o teste tentou acessar o Google de verdade")

    rede_original = (nuvem._obter_cliente, sheets_sync._obter_cliente)
    nuvem._obter_cliente = sheets_sync._obter_cliente = _sem_rede
    datetime_original = v.datetime
    pasta = Path(tempfile.mkdtemp(prefix="_smoke_pendencias_"))
    fx.isolar_preferencias(pasta)
    caminho = fx.criar(pasta)
    fx.apontar_modulos_para(caminho)
    sessao_mod.iniciar(ADMIN)

    def hash_do_arquivo() -> str:
        return hashlib.sha256(caminho.read_bytes()).hexdigest()

    def venda_(id_venda):
        return bd.ler_vendas(caminho).set_index("ID_VENDA").loc[id_venda]

    def resumo(id_venda):
        return next(r for r in v.listar_vendas() if r.id == id_venda)

    def historico_da_venda(id_venda):
        h = bd.ler_historico(caminho)
        return h[(h["ID"] == id_venda) & (h["TIPO"] == v.TIPO_PENDENCIA)]

    def com_banco_fechado(cpf, equipamento, banco="Banco Exemplo"):
        id_venda, id_proposta = v.criar_venda(cpf, [equipamento], _pedido(banco, EQUIPAMENTO=equipamento))
        v.mudar_status_proposta(id_proposta, APROVADO)
        v.fechar_proposta(id_proposta)
        return id_venda, id_proposta

    try:
        linha("1) Varias pendencias na mesma venda: texto, datas, ordem, ids")
        assert v.TIPO_PENDENCIA == "pendencia"
        id_v1, id_p1 = v.criar_venda(fx.CPF_MARIA, ["Cadeira Exemplo"], _pedido("Banco Exemplo", EQUIPAMENTO="Cadeira Exemplo"))
        assert resumo(id_v1).pendencias == () and resumo(id_v1).pendencias_abertas == 0
        assert v.pendencias_da_venda(id_v1) == [] and venda_(id_v1)["PENDENCIAS"] == ""
        v.datetime = _Relogio
        _Relogio.agora_combinado = datetime(2026, 3, 2, 10, 30, 15, 999)
        id_a = v.adicionar_pendencia(id_v1, "  Equipamento ainda não chegou  ")
        _Relogio.agora_combinado = datetime(2026, 3, 5, 9, 0, 0)
        id_b = v.adicionar_pendencia(id_v1, 'Falta "comprovante" de renda, ação urgente; 50% = ok')
        id_c = v.adicionar_pendencia(id_v1, "Cliente vai trazer o avalista")
        assert len({id_a, id_b, id_c}) == 3 and all(i.startswith("PD-") for i in (id_a, id_b, id_c))
        lista = v.pendencias_da_venda(id_v1)
        assert [p.id for p in lista] == [id_a, id_b, id_c], "na ordem em que foram criadas"
        assert lista[0].texto == "Equipamento ainda não chegou", "o texto vem sem os espaços das pontas"
        assert lista[1].texto == 'Falta "comprovante" de renda, ação urgente; 50% = ok'
        assert lista[0].criada == datetime(2026, 3, 2, 10, 30, 15) and lista[0].resolvida is None and lista[0].aberta
        assert lista[1].criada == datetime(2026, 3, 5, 9, 0, 0)
        r = resumo(id_v1)
        assert r.pendencias == tuple(lista) and r.pendencias_abertas == 3
        print("OK: três pendências, na ordem, com texto (aspas, vírgula, acento, ;, %, =) e data de criação.")

        linha("2) Texto vazio e recusado; nada e gravado")
        antes = hash_do_arquivo()
        for vazio in ("", "   ", "\n\t", None):
            _recusa("Diga o que está pendente", v.adicionar_pendencia, id_v1, vazio)
        _recusa("Venda não encontrada", v.adicionar_pendencia, "V-nao-existe", "x")
        assert hash_do_arquivo() == antes and len(v.pendencias_da_venda(id_v1)) == 3
        print("OK: texto vazio (ou só espaços) e venda inexistente recusam com mensagem, sem gravar.")

        linha("3) Resolver uma: as outras seguem abertas; resolver de novo, id errado e venda errada recusam")
        _Relogio.agora_combinado = datetime(2026, 3, 20, 16, 45, 0)
        v.resolver_pendencia(id_v1, id_a)
        lista = v.pendencias_da_venda(id_v1)
        assert [p.aberta for p in lista] == [False, True, True]
        assert lista[0].resolvida == datetime(2026, 3, 20, 16, 45, 0) and lista[0].criada == datetime(2026, 3, 2, 10, 30, 15)
        assert lista[1].resolvida is None and lista[0].texto == "Equipamento ainda não chegou"
        assert resumo(id_v1).pendencias_abertas == 2, "a resolvida fica na lista mas não conta como aberta"
        antes = hash_do_arquivo()
        _recusa("já está resolvida", v.resolver_pendencia, id_v1, id_a)
        _recusa("Pendência não encontrada", v.resolver_pendencia, id_v1, "PD-inexistente")
        _recusa("Pendência não encontrada", v.resolver_pendencia, id_v1, "")
        _recusa("Venda não encontrada", v.resolver_pendencia, "V-nao-existe", id_a)
        assert hash_do_arquivo() == antes, "as recusas não gravam nada"
        v.resolver_pendencia(id_v1, id_b)
        v.resolver_pendencia(id_v1, id_c)
        assert resumo(id_v1).pendencias_abertas == 0 and len(resumo(id_v1).pendencias) == 3
        print("OK: resolver guarda a data; a resolvida continua na lista; resolver de novo/errado recusa sem gravar.")

        linha("4) Cada criacao e resolucao vai para o HISTORICO (TIPO pendencia, ID = a venda)")
        h = historico_da_venda(id_v1)
        assert list(h["PARA"]) == ["Adicionada"] * 3 + ["Resolvida"] * 3, list(h["PARA"])
        assert list(h["DE"]) == [""] * 3 + ["Adicionada"] * 3
        assert list(h["MOTIVO"])[:3] == [p.texto for p in v.pendencias_da_venda(id_v1)]
        assert list(h["MOTIVO"])[3:] == list(h["MOTIVO"])[:3], "a resolução guarda o texto resolvido"
        assert h["QUANDO"].iloc[0] == pd.Timestamp("2026-03-02 10:30:15") and h["QUANDO"].iloc[3] == pd.Timestamp("2026-03-20 16:45:00")
        # nao e troca de status: o "voltar" da proposta nao enxerga essas linhas
        v.mudar_status_proposta(id_p1, APROVADO)
        assert v._anterior(id_p1) == propostas_mod.STATUS_EM_ANALISE
        assert v._pilhas_de_anteriores(bd.ler_historico(caminho)).get(id_v1) is None
        print("OK: 3 'Adicionada' + 3 'Resolvida', com o texto e a data; o voltar das propostas ignora.")

        linha("5) So o administrador escreve")
        sessao_mod.encerrar()
        antes = hash_do_arquivo()
        for chamada in (lambda: v.adicionar_pendencia(id_v1, "x"), lambda: v.resolver_pendencia(id_v1, id_a)):
            try:
                chamada()
                raise AssertionError("sem sessão não pode gravar")
            except sessao_mod.PermissaoNegada:
                pass
        assert hash_do_arquivo() == antes
        sessao_mod.iniciar(ADMIN)
        print("OK: sem sessão de administrador, nada grava.")

        linha("6) Em qualquer etapa; o status da venda nao muda; nada trava")
        v.datetime = datetime_original
        # Aguardando bancos
        id_v2, id_p2 = v.criar_venda(fx.CPF_JOAO, ["Laser Exemplo"], _pedido("Banco Exemplo", EQUIPAMENTO="Laser Exemplo"))
        # Em andamento (um banco fechado)
        id_v3, id_p3 = com_banco_fechado(fx.CPF_ANA, "Mesa Exemplo")
        # Efetivada
        id_v4, id_p4 = com_banco_fechado(fx.CPF_ANA, "Autoclave Exemplo")
        v.efetivar_proposta(id_p4)
        # Perdida
        id_v5, _ = v.criar_venda(fx.CPF_MARIA, ["Torno Exemplo"], _pedido("Banco Exemplo", EQUIPAMENTO="Torno Exemplo"))
        v.mudar_status_venda(id_v5, v.VENDA_PERDIDA, motivo="Desistiu")
        esperados = {id_v2: v.VENDA_AGUARDANDO, id_v3: v.VENDA_EM_ANDAMENTO, id_v4: v.VENDA_EFETIVADA, id_v5: v.VENDA_PERDIDA}
        for id_venda, status in esperados.items():
            assert resumo(id_venda).status == status
            v.adicionar_pendencia(id_venda, "Equipamento ainda não chegou")
            assert resumo(id_venda).status == status and venda_(id_venda)["STATUS"] == status, "a pendência não muda o status"
            assert resumo(id_venda).pendencias_abertas == 1
        # com a pendencia aberta, nada do que o Pendente de entrega travava trava
        outro = v.mandar_a_outro_banco(id_v4, _pedido("Banco Segundo", EQUIPAMENTO="Outro Item"))
        v.mudar_status_proposta(outro, APROVADO)
        v.fechar_proposta(outro)
        assert venda_(id_v4)["STATUS"] == v.VENDA_EM_ANDAMENTO, "fechar outro banco reabre o andamento, com a pendência aberta"
        v.mudar_status_proposta(outro, ASSINADO)
        v.mudar_status_proposta(id_p3, ASSINADO)
        v.voltar_proposta(id_p3)
        v.mudar_status_venda(id_v3, v.VENDA_NAO_EFETIVADA, motivo="Desistiu")
        assert v.voltar_venda(id_v3) == v.VENDA_EM_ANDAMENTO
        assert resumo(id_v3).pendencias_abertas == 1 and resumo(id_v4).pendencias_abertas == 1
        # o status continua sendo so o dos bancos/desfechos
        assert v.STATUS_DA_VENDA == [v.VENDA_AGUARDANDO, v.VENDA_EM_ANDAMENTO, v.VENDA_EFETIVADA, v.VENDA_NAO_EFETIVADA, v.VENDA_PERDIDA]
        print("OK: pendência em Aguardando, Em andamento, Efetivada e Perdida; status intacto; mandar a outro banco, fechar, "
              "mudar etapa e desfechos seguem livres.")

        linha("7) A pendencia nunca conta para a parada de 7 dias")
        limite = propostas_mod.DIAS_PROPOSTA_PARADA
        hoje = date.today()
        v.datetime = _Relogio
        # venda Aguardando bancos criada ha 20 dias, sem proposta: parada de verdade; a pendencia e de 3 dias atras
        _Relogio.agora_combinado = datetime.combine(hoje - timedelta(days=20), datetime.min.time()) + timedelta(hours=9)
        id_v6, _ = v.criar_venda(fx.CPF_JOAO, ["Item Parado"])
        _Relogio.agora_combinado = datetime.combine(hoje - timedelta(days=3), datetime.min.time()) + timedelta(hours=9)
        id_pd6 = v.adicionar_pendencia(id_v6, "Falta documento")
        # venda Efetivada ha 20 dias, com pendencia aberta desde entao
        _Relogio.agora_combinado = datetime.combine(hoje - timedelta(days=20), datetime.min.time()) + timedelta(hours=9)
        id_v7, id_p7 = com_banco_fechado(fx.CPF_ANA, "Item Efetivado")
        v.efetivar_proposta(id_p7)
        v.adicionar_pendencia(id_v7, "Equipamento ainda não chegou")
        # venda Em andamento: banco aprovado e fechado ha 2 dias; pendencia aberta ha 30 dias nao a deixa parada nem alerta
        _Relogio.agora_combinado = datetime.combine(hoje - timedelta(days=30), datetime.min.time()) + timedelta(hours=9)
        id_v8, id_p8 = v.criar_venda(fx.CPF_MARIA, ["Item Andando"], _pedido("Banco Exemplo", EQUIPAMENTO="Item Andando"))
        v.adicionar_pendencia(id_v8, "Cliente viaja")
        _Relogio.agora_combinado = datetime.combine(hoje - timedelta(days=2), datetime.min.time()) + timedelta(hours=9)
        v.mudar_status_proposta(id_p8, APROVADO)
        v.fechar_proposta(id_p8)
        v.datetime = datetime_original

        alertas = {(a.tipo, a.id): a for a in v.alertas(hoje)}
        # a venda parada alerta pelos 20 dias dela - a pendencia de 3 dias atras nao zera a contagem
        assert alertas[(v.TIPO_VENDA, id_v6)].dias == 20, alertas[(v.TIPO_VENDA, id_v6)]
        assert resumo(id_v6).parada and resumo(id_v6).dias_na_etapa == 20
        assert (v.TIPO_VENDA, id_v7) not in alertas and (v.TIPO_PROPOSTA, id_p7) not in alertas, "Efetivada com pendência: sem alerta"
        assert not resumo(id_v7).parada, "sem selo de parada por causa da pendência"
        assert (v.TIPO_VENDA, id_v8) not in alertas and (v.TIPO_PROPOSTA, id_p8) not in alertas, "banco parado há 2 dias: sem alerta"
        assert not resumo(id_v8).parada and resumo(id_v8).pendencias_abertas == 1
        for daqui_a in (limite, 60):  # e nada muda com o tempo: a Efetivada/Em andamento so alerta pela etapa do banco
            em = {(a.tipo, a.id) for a in v.alertas(hoje + timedelta(days=daqui_a))}
            assert (v.TIPO_VENDA, id_v7) not in em and (v.TIPO_PROPOSTA, id_p7) not in em
            assert not any("pendência" in a.texto.lower() or "pendencia" in a.texto.lower()
                           for a in v.alertas(hoje + timedelta(days=daqui_a)))
        assert (v.TIPO_PROPOSTA, id_p8) in {(a.tipo, a.id) for a in v.alertas(hoje + timedelta(days=limite))}, \
            "o banco fechado parado alerta pela etapa dele (como na 1.2-A)"
        # resolver a pendencia tambem nao mexe na contagem da venda parada
        v.resolver_pendencia(id_v6, id_pd6)
        assert {(a.tipo, a.id): a.dias for a in v.alertas(hoje)}[(v.TIPO_VENDA, id_v6)] == 20
        print("OK: a pendência (criada ou resolvida) não gera alerta, não dá selo de parada e não zera a parada da venda.")

        linha("8) Ida e volta com a nuvem simulada")
        locais = {bd.ABA_PROPOSTAS: bd.ler_propostas(caminho), bd.ABA_VENDAS: bd.ler_vendas(caminho),
                  bd.ABA_HISTORICO: bd.ler_historico(caminho)}
        original_bruta = nuvem._ler_aba_bruta

        def _aba_da_nuvem(nome):  # o que o Google devolveria: as linhas exatamente como foram mandadas
            linhas = sheets_sync._preparar_linhas(locais[nome])
            return pd.DataFrame(linhas[1:], columns=linhas[0]) if len(linhas) > 1 else pd.DataFrame()

        nuvem._ler_aba_bruta = _aba_da_nuvem
        try:
            da_nuvem = {bd.ABA_PROPOSTAS: nuvem.ler_propostas(), bd.ABA_VENDAS: nuvem.ler_vendas(),
                        bd.ABA_HISTORICO: nuvem.ler_historico()}
        finally:
            nuvem._ler_aba_bruta = original_bruta
        nuvem_v, locais_v = da_nuvem[bd.ABA_VENDAS], locais[bd.ABA_VENDAS]
        assert "PENDENCIAS" in bd.VENDAS_COLUNAS and bd.VENDAS_COLUNAS[-1] == "PENDENCIAS", "a coluna nova vai no fim"
        assert list(nuvem_v["PENDENCIAS"]) == list(locais_v["PENDENCIAS"]) and any(locais_v["PENDENCIAS"])
        assert list(da_nuvem[bd.ABA_HISTORICO]["TIPO"]) == list(locais[bd.ABA_HISTORICO]["TIPO"])
        antes = {r.id: r.pendencias for r in v.listar_vendas()}
        bd.escrever_tudo(caminho, bd.ler_clientes(caminho), bd.ler_equipamentos(caminho), bd.ler_vendedores(caminho),
                         da_nuvem[bd.ABA_PROPOSTAS], vendas=nuvem_v, historico=da_nuvem[bd.ABA_HISTORICO])
        depois = {r.id: r.pendencias for r in v.listar_vendas()}
        assert depois == antes and sum(len(p) for p in depois.values()) >= 9, "nenhuma pendência (nem data) se perde na volta"
        assert len(depois[id_v1]) == 3 and all(not p.aberta for p in depois[id_v1])
        assert depois[id_v1][1].texto == 'Falta "comprovante" de renda, ação urgente; 50% = ok'
        print("OK: a coluna PENDENCIAS (texto, datas, resolvidas) vai e volta pela nuvem simulada sem perder nada.")

        linha("9) Planilha de antes da coluna PENDENCIAS abre sem erro e ganha a coluna ao gravar")
        for ws_nome in (bd.ABA_VENDAS,):
            wb = openpyxl.load_workbook(caminho)
            ws = wb[ws_nome]
            assert ws.cell(row=1, column=9).value == "PENDENCIAS"
            ws.delete_cols(9)
            wb.save(caminho)
            wb.close()
        sem_coluna = bd.ler_vendas(caminho)
        assert list(sem_coluna.columns) == bd.VENDAS_COLUNAS and set(sem_coluna["PENDENCIAS"]) == {""}, "lê como vazio"
        r = resumo(id_v1)
        assert r.pendencias == () and r.pendencias_abertas == 0 and len(v.listar_vendas()) >= 8
        assert v.pendencias_da_venda(id_v1) == [] and v.alertas(hoje) is not None
        assert resumo_vendas.texto_da_venda(id_v1, hoje)  # o resumo da venda tambem abre
        wb = openpyxl.load_workbook(caminho)
        assert wb[bd.ABA_VENDAS].cell(row=1, column=9).value is None
        wb.close()
        v.editar_venda(id_v2, ["Laser Exemplo"], "obs depois de abrir sem a coluna")  # uma gravacao qualquer
        wb = openpyxl.load_workbook(caminho)
        assert wb[bd.ABA_VENDAS].cell(row=1, column=9).value == "PENDENCIAS", "o cabeçalho voltou na primeira gravação"
        wb.close()
        id_nova = v.adicionar_pendencia(id_v2, "Primeira pendência da planilha antiga")
        assert [p.id for p in v.pendencias_da_venda(id_v2)][-1] == id_nova
        assert venda_(id_v1)["PENDENCIAS"] == "" and venda_(id_v2)["OBSERVAÇÕES"] == "obs depois de abrir sem a coluna"
        print("OK: sem a coluna, lê vazio e tudo abre; a primeira gravação põe o cabeçalho e a pendência nova grava.")

        linha("10) Celula ilegivel (editada a mao): so as pendencias DESSA venda ficam bloqueadas, com mensagem; o resto abre")
        antes_das_outras = {r.id: r.pendencias for r in v.listar_vendas() if r.id != id_v2}
        id_pd2 = v.adicionar_pendencia(id_v2, "uma pendência que vai ficar ilegível")
        vendas = bd.ler_vendas(caminho)
        vendas.loc[vendas["ID_VENDA"] == id_v2, "PENDENCIAS"] = "isto não é uma lista"
        bd.escrever_vendas_e_propostas(caminho, vendas=vendas)
        # a lista de vendas (e com ela a Ficha, Todas as Propostas e os resumos) NAO cai: a venda vem sem pendencias, com o erro
        todas = {r.id: r for r in v.listar_vendas()}
        assert id_v2 in todas and todas[id_v2].pendencias == () and todas[id_v2].pendencias_abertas == 0
        msg = todas[id_v2].erro_pendencias
        assert "PENDENCIAS" in msg and id_v2 in msg and "ilegível" in msg, msg
        assert "apague" not in msg.lower() and "avise o suporte" in msg and "nuvem" in msg and "backups" in msg, msg
        assert all(x.erro_pendencias == "" for i, x in todas.items() if i != id_v2), "o erro e so dessa venda"
        assert {i: x.pendencias for i, x in todas.items() if i != id_v2} == antes_das_outras, "as outras venda seguem intactas"
        assert resumo_vendas.texto_da_venda(id_v2, date.today()) and v.alertas(date.today()) is not None
        assert resumo_vendas.montar_resumo(resumo_vendas.periodo(resumo_vendas.SEMANA, resumo_vendas.ATUAL), None).total >= 1
        # so as gravacoes de pendencia dessa venda ficam bloqueadas, com a mesma mensagem (e nada e gravado)
        antes = hash_do_arquivo()
        for chamada in (lambda: v.pendencias_da_venda(id_v2), lambda: v.adicionar_pendencia(id_v2, "x"),
                        lambda: v.resolver_pendencia(id_v2, id_pd2)):
            try:
                chamada()
                raise AssertionError("deveria recusar a célula ilegível")
            except v.ErroVenda as exc:
                assert "PENDENCIAS" in str(exc) and id_v2 in str(exc) and "ilegível" in str(exc), str(exc)
                assert "apague" not in str(exc).lower() and "suporte" in str(exc), str(exc)
        assert hash_do_arquivo() == antes
        v.adicionar_pendencia(id_v1, "as outras vendas aceitam pendência normalmente")
        # uma gravacao qualquer da venda ilegivel NAO perde nem mexe na celula (o dado fica para quem for corrigir)
        v.editar_venda(id_v2, ["Laser Exemplo"], "outra observação")
        assert venda_(id_v2)["PENDENCIAS"] == "isto não é uma lista"
        for ruim in ('[{"id":"PD-1"}]', '{"id":"PD-1"}', '[1,2]', '[{"id":"PD-1","texto":"x","criada":"ontem","resolvida":""}]'):
            _recusa("ilegível", v.pendencias_do_texto, ruim, "V-x")
        vendas = bd.ler_vendas(caminho)
        vendas.loc[vendas["ID_VENDA"] == id_v2, "PENDENCIAS"] = ""
        bd.escrever_vendas_e_propostas(caminho, vendas=vendas)
        assert resumo(id_v2).pendencias == () and resumo(id_v2).erro_pendencias == ""
        assert v.pendencias_do_texto("") == () and v.pendencias_do_texto(None) == () and v.pendencias_do_texto("  ") == ()
        print("OK: célula ilegível: a lista, a Ficha e os resumos abrem; só as pendências dessa venda bloqueiam, com mensagem que "
              "nunca manda apagar; a célula não é tocada; vazia lê como sem pendências.")

        linha("11) NF enviada (nome novo) e Nota fiscal enviada (nome antigo, sinonimo)")
        assert propostas_mod.STATUS_NOTA_ENVIADA == "NF enviada" and NOTA == "NF enviada"
        assert propostas_mod.STATUS_NOTA_ENVIADA in propostas_mod.STATUS_OFICIAIS
        assert "Nota fiscal enviada" not in propostas_mod.STATUS_OFICIAIS, "o nome antigo não é mais oficial"
        assert v.ETAPAS_DO_FECHADO == [APROVADO, ASSINADO, "NF enviada", EFETIVADO]
        for status in ("NF enviada", "nf enviada", "NF ENVIADA", " nf enviada ", "Nota fiscal enviada", "NOTA FISCAL ENVIADA"):
            assert propostas_mod.etapa_status(status) == propostas_mod.ETAPA_APROVADO, status
            assert propostas_mod.eh_aprovado_nao_efetivado(status), status
        for antigo in ("Nota fiscal enviada", "NOTA FISCAL ENVIADA", "Nota Fiscal Anexada", "NF enviada", "nf Enviada"):
            assert migracao_legado.converter_status(antigo) == "NF enviada", antigo
        # uma proposta fechada ainda gravada com o nome antigo: a venda segue Em andamento e a planilha abre
        id_v9, id_p9 = com_banco_fechado(fx.CPF_ANA, "Item Nome Antigo")
        propostas, _ = v._ler()
        propostas.loc[propostas["ID_PROPOSTA"] == id_p9, "STATUS"] = "Nota fiscal enviada"
        bd.escrever_vendas_e_propostas(caminho, propostas=propostas)
        assert resumo(id_v9).status == v.VENDA_EM_ANDAMENTO
        v.efetivar_proposta(id_p9)
        assert resumo(id_v9).status == v.VENDA_EFETIVADA
        print("OK: 'NF enviada' é o oficial; 'Nota fiscal enviada' ainda é lida como a mesma etapa e a conversão a troca.")

        linha("12) O bloco Pendencias do resumo aparece so com pendencias abertas (sem mais nada no bloco)")
        per = resumo_vendas.periodo(resumo_vendas.SEMANA, resumo_vendas.ATUAL, date(2026, 10, 21))
        so_pendencias = resumo_vendas.ResumoDoVendedor(per, None, 1, (), (), (), (), (),
                                                       com_pendencias=(("Cliente Exemplo", ("Falta o contrato", "Obra atrasada")),))
        assert so_pendencias.tem_pendencias
        texto = resumo_vendas.texto_whatsapp(so_pendencias)
        assert texto.endswith("*Pendências*\nPendências em aberto:\n• Cliente Exemplo: Falta o contrato\n• Cliente Exemplo: Obra atrasada"), texto
        sem = resumo_vendas.ResumoDoVendedor(per, None, 1, (), (), (), (), ())
        assert not sem.tem_pendencias and "Pendências" not in resumo_vendas.texto_whatsapp(sem)
        print("OK: só com pendência aberta o bloco aparece; sem nada, não.")

        linha("TUDO OK")
    finally:
        v.datetime = datetime_original
        nuvem._obter_cliente, sheets_sync._obter_cliente = rede_original
        fx.restaurar_modulos()
        sessao_mod.encerrar()


if __name__ == "__main__":
    main()
