"""Testa o cadastro de Equipamentos (core/equipamentos.py e desktop/widgets/cadastro_equipamentos.py): cadastrar,
editar (o nome novo vai para propostas e vendas), desativar (sai das listas), excluir so o que ninguem usa, a
planilha de antes da coluna ATIVO, e "trazer para o cadastro" os nomes ja usados - os simples entram, os que podem
ser dois equipamentos num nome so vao para o usuario decidir, com backup antes.

Planilha 100% ficticia em pasta temporaria, preferencias isoladas e o Google desligado (qualquer acesso derruba).

Rodar com: venv/Scripts/python.exe scripts/smoke_test_equipamentos_cadastro.py
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl
from PySide6.QtWidgets import QApplication, QDialog

import config

config.SINCRONIZACAO_GOOGLE_ATIVADA = False
import fixture_ficticia as fx
from ambiente_de_teste import Mensagens
from core import bancos as bancos_mod
from core import data_store as bd
from core import data_store_sheets as nuvem
from core import equipamentos as eq
from core import propostas as propostas_mod
from core import sessao as sessao_mod
from core import sheets_sync
from core import vendas as v
from desktop.widgets.campos_da_venda import FormularioDaVenda
from desktop.theme import TEMA_ESCURO, build_stylesheet
from desktop.widgets.cadastro_equipamentos import CadastroEquipamentos, TrazerEquipamentosDialog

ADMIN = sessao_mod.Sessao(papel=sessao_mod.PAPEL_ADMIN, nome_usuario="Administrador")


def linha(titulo: str) -> None:
    print(f"\n{'=' * 60}\n{titulo}\n{'=' * 60}")


def _recusa(trecho: str, funcao, *args) -> None:
    try:
        funcao(*args)
    except eq.ErroEquipamento as exc:
        assert trecho in str(exc), (trecho, str(exc))
        return
    raise AssertionError(f"deveria recusar ({trecho})")


def main() -> None:
    def _sem_rede(*_a, **_k):
        raise AssertionError("o teste tentou acessar o Google de verdade")

    rede_original = (nuvem._obter_cliente, sheets_sync._obter_cliente)
    nuvem._obter_cliente = sheets_sync._obter_cliente = _sem_rede
    pasta = Path(tempfile.mkdtemp(prefix="_smoke_equip_"))
    registro_antes = fx.instantaneo_do_registro()
    fx.isolar_preferencias(pasta)
    caminho = fx.criar(pasta)
    fx.apontar_modulos_para(caminho)
    sessao_mod.iniciar(ADMIN)
    app = QApplication.instance() or QApplication(sys.argv)
    app.setStyleSheet(build_stylesheet(TEMA_ESCURO))
    try:
        with Mensagens() as msgs:
            banco = bancos_mod.nomes_ativos()[0]

            linha("1) Planilha de antes da coluna ATIVO: lê todos ativos; a 1ª gravação põe o cabeçalho")
            with open(caminho, "rb") as f:
                assert openpyxl.load_workbook(f)[bd.ABA_EQUIPAMENTOS].cell(row=1, column=7).value is None
            iniciais = eq.nomes_ativos()
            assert iniciais == ["Equipamento Modelo X", "Equipamento Modelo Y"], iniciais
            eq.cadastrar({"EQUIPAMENTO": "  Laser   Exemplo ", "FORNECEDOR": "Fornecedor Exemplo", "PARCELAS": 48,
                          "VALOR LÍQUIDO/REFERÊNCIA (R$)": 80000})
            with open(caminho, "rb") as f:
                assert openpyxl.load_workbook(f)[bd.ABA_EQUIPAMENTOS].cell(row=1, column=7).value == "ATIVO"
            assert eq.nomes_ativos() == ["Equipamento Modelo X", "Equipamento Modelo Y", "Laser Exemplo"]
            print("OK: sem a coluna, todos ativos; cadastrar põe o cabeçalho 'ATIVO' e o nome sem espaços sobrando.")

            linha("2) Regras: nome obrigatório, sem repetir, sem '+'; números válidos")
            _recusa("obrigatório", eq.cadastrar, {"EQUIPAMENTO": "  "})
            _recusa("já está no cadastro", eq.cadastrar, {"EQUIPAMENTO": "laser exemplo"})
            _recusa('não pode ter "+"', eq.cadastrar, {"EQUIPAMENTO": "Laser + Cadeira"})
            _recusa("maior que zero", eq.cadastrar, {"EQUIPAMENTO": "Cadeira Exemplo", "PARCELAS": 0})
            eq.cadastrar({"EQUIPAMENTO": "Cadeira Exemplo"})
            print("OK: nome em branco, repetido (sem diferenciar maiúsculas) ou com '+' é recusado; parcelas > 0.")

            linha("3) Renomear leva o nome novo para as propostas e vendas; desativar tira das listas")
            id_venda, id_p = v.criar_venda(fx.CPF_MARIA, ["Laser Exemplo", "Cadeira Exemplo"],
                                           {"BANCO": banco, "VALOR (R$)": 1000, "MESES": 12})
            v.mandar_a_outro_banco(id_venda, {"BANCO": banco, "VALOR (R$)": 500, "MESES": 12, "EQUIPAMENTO": "Laser Exemplo"})
            assert eq.salvar("Laser Exemplo", {"EQUIPAMENTO": "Laser Exemplo 3000", "FORNECEDOR": "Outro"}) == 2
            venda = bd.ler_vendas(caminho).set_index("ID_VENDA").loc[id_venda]
            assert venda["EQUIPAMENTOS"] == "Laser Exemplo 3000 + Cadeira Exemplo"
            assert list(v.propostas_da_venda(id_venda)["EQUIPAMENTO"]) == [
                "Laser Exemplo 3000 + Cadeira Exemplo", "Laser Exemplo 3000"]
            cadastro = eq.listar_cadastro().set_index("EQUIPAMENTO")
            assert cadastro.loc["Laser Exemplo 3000", "FORNECEDOR"] == "Outro" and cadastro.loc["Laser Exemplo 3000", "PROPOSTAS"] == 2
            _recusa("Já existe outro", eq.salvar, "Cadeira Exemplo", {"EQUIPAMENTO": "laser exemplo 3000"})
            eq.definir_ativo("Cadeira Exemplo", False)
            assert "Cadeira Exemplo" not in eq.nomes_ativos()
            d = FormularioDaVenda(venda=next(r for r in v.listar_vendas() if r.id == id_venda))
            assert [c.currentText() for c in d.campos.equipamentos.linhas] == ["Laser Exemplo 3000", "Cadeira Exemplo"], \
                "desativado: a venda que já o tem continua mostrando"
            nova = FormularioDaVenda()
            nova.campos.equipamentos.mais.click()
            linhas = nova.campos.equipamentos.linhas
            opcoes = [linhas[1].itemText(i) for i in range(linhas[1].count())]
            assert "Cadeira Exemplo" not in opcoes and "Laser Exemplo 3000" in opcoes and not linhas[1].isEditable()
            eq.definir_ativo("Cadeira Exemplo", True)
            print("OK: renomear muda propostas e vendas (2 propostas); desativado sai da lista nova mas fica onde já está.")

            linha("4) Excluir só o que nenhuma proposta usa")
            _recusa("não pode ser excluído", eq.excluir, "Laser Exemplo 3000")
            eq.cadastrar({"EQUIPAMENTO": "Autoclave Exemplo"})
            eq.excluir("Autoclave Exemplo")
            assert "Autoclave Exemplo" not in eq.nomes_ativos()
            print("OK: usado em proposta não se exclui (desativa); sem uso, sai do cadastro.")

            linha("5) Trazer os nomes já usados: simples entram; os 'juntos' o usuário decide; backup antes")
            propostas_mod.adicionar_proposta({"CPF": fx.CPF_JOAO, "VALOR (R$)": 1000, "BANCO": banco,
                                              "EQUIPAMENTO": "Mesa Antiga Exemplo"})
            propostas_mod.adicionar_proposta({"CPF": fx.CPF_ANA, "VALOR (R$)": 1000, "BANCO": banco,
                                              "EQUIPAMENTO": "Laser Velho E Cadeira Velha"})
            propostas_mod.adicionar_proposta({"CPF": fx.CPF_ANA, "VALOR (R$)": 1000, "BANCO": banco,
                                              "EQUIPAMENTO": "Ultrassom E Doppler"})
            simples, juntos = eq.nomes_fora_do_cadastro()
            assert "Mesa Antiga Exemplo" in simples and juntos == ["Laser Velho E Cadeira Velha", "Ultrassom E Doppler"]
            _recusa("Decida antes", eq.trazer_para_o_cadastro, simples, {})
            cadastro = CadastroEquipamentos()
            assert not cadastro.aviso_fora.isHidden() and "2 para você decidir" in cadastro.texto_fora.text()
            dialogo = TrazerEquipamentosDialog(simples, juntos)
            assert not dialogo.botao_trazer.isEnabled(), "nada marcado de início: não deixa trazer"
            grupo1, partes1 = dialogo.respostas["Laser Velho E Cadeira Velha"]
            assert partes1.text() == "Laser Velho, Cadeira Velha" and not partes1.isEnabled()
            grupo1.button(2).click()
            assert not dialogo.botao_trazer.isEnabled()
            dialogo.respostas["Ultrassom E Doppler"][0].button(1).click()  # é um aparelho só
            assert dialogo.botao_trazer.isEnabled()
            dialogo._trazer()
            assert dialogo.result() == QDialog.DialogCode.Accepted and "entraram no cadastro" in msgs.ultima()[2]
            r = dialogo.resultado
            assert r["propostas_alteradas"] == 1 and Path(r["backup"]).exists() and "pre-equipamentos" in Path(r["backup"]).name
            ativos = eq.nomes_ativos()
            for nome in ("Mesa Antiga Exemplo", "Laser Velho", "Cadeira Velha", "Ultrassom E Doppler"):
                assert nome in ativos, nome
            assert "Laser Velho E Cadeira Velha" not in ativos
            textos = set(bd.ler_propostas(caminho)["EQUIPAMENTO"])
            assert "Laser Velho + Cadeira Velha" in textos and "Ultrassom E Doppler" in textos
            assert eq.nomes_fora_do_cadastro() == ([], []), "depois de trazer, nada falta"
            cadastro.recarregar()
            assert cadastro.aviso_fora.isHidden(), "nada faltando: o aviso some"
            print("OK: simples entram sozinhos; 'X E Y' separado vira 'X + Y' nas propostas; 'um só' fica como está; "
                  "backup 'pre-equipamentos'; o aviso some.")

            linha("6) A seção Equipamentos em Cadastros")
            titulos = [c.item.titulo for c in cadastro.cards]
            assert "Laser Exemplo 3000" in titulos and "Mesa Antiga Exemplo" in titulos
            assert not cadastro.cards[0].no_lugar, "vários campos: o card abre para baixo"
            card = next(c for c in cadastro.cards if c.item.titulo == "Laser Exemplo 3000")
            assert "Outro" in card.item.detalhe and "2 propostas" in card.item.detalhe
            assert card.item.valores["parcelas"] == "48" and card.item.valores["referencia"] == "80000"
            print("OK: um card por equipamento, abre para baixo (vários campos), com fornecedor e uso no detalhe.")
        linha("TUDO OK")
    finally:
        nuvem._obter_cliente, sheets_sync._obter_cliente = rede_original
        fx.restaurar_modulos()
        sessao_mod.encerrar()
        assert fx.instantaneo_do_registro() == registro_antes, "o teste mexeu nas preferências reais"


if __name__ == "__main__":
    main()
